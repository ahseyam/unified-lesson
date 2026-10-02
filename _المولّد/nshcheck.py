# -*- coding: utf-8 -*-
"""فحصُ نشرات المنصة — المحتوى من الـDOCX والشكلُ من الـPDF.

⛔ لا يُقاس المحتوى من نصِّ الـPDF: خطُّ الجزيرة يربط حروفَه، وجدولُ ToUnicode
   فيه يُخرج «کل» بكافٍ فارسية و«'ي» بدل «في» و«-.» بدل «تر». فالمطابقةُ
   الحرفيةُ تفشل على مستندٍ سليم. المحتوى يُقرأ من الـDOCX حيث النصُّ حقيقي.

⛔ ولا يُقاس الشكلُ من الـDOCX: عددُ الصفحات وموضعُ النصِّ والكليشةُ لا تُعرف
   إلا بعد الإخراج. فالشكلُ من الـPDF بالبكسل.

الاستعمال: python3 nshcheck.py <الملف.docx>
"""
import os
import sys

import fitz
from docx import Document

MM = 72 / 25.4
BAND = 24 * MM            # شريطُ الكليشة — عنوانُ المستند فيه قصداً
# ⚠️ هامشُ الصفحة ١٠ مم، وصندوقُ الكتلة الذي يُبلّغ عنه PyMuPDF يتّسع عليه
#    مليمترين فيبدأ عند ٨. فحدُّ الإنذار ٧ مم — أضيقُ منه تجاوزٌ حقيقي.
SIDE = 7 * MM
FOOT = 24 * MM


def docx_text(p):
    d = Document(p)
    out = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                out.append(c.text)
    return "\n".join(out)


# ⛔ **وكانت النشراتُ الإنجليزيةُ بلا فحصٍ أصلاً**: الشواهدُ كلُّها عربيةٌ
#    («قبل أن تبدأ» …) فلم تُقَس عليها سبعُ نشراتٍ منشورةٍ — والشكلُ والتسرّبُ
#    لا لغةَ لهما. فصار المحتوى بقائمةٍ لكلِّ لغة، والشكلُ لهما معاً.
#    (أمسكه مسحُ ما بقي ١ أكتوبر ٢٠٢٦)
NEED_AR = [("الرابط", "ahseyam.github.io/unified-lesson"),
           ("التوقيع", "إدارةالتخطيطوالاعتمادالمدرسي"),
           ("قبل أن تبدأ", "قبلأنتبد"),
           ("خطوةً خطوة", "خطوةًخطوة"),
           ("ما يقع فيه الخطأ", "يقعفيهالخطأ"),
           ("أسئلةٌ متكرّرة", "أسئلةٌمتكرّرة"),
           ("بمن تتصل", "بمنتتصل"),
           ("شاهدُ اللقطة", "لقطةٌمنالمنصةنفسِها")]
# ⚠️ وهذه مقروءةٌ من `nshbuild_en.py` حرفاً حرفاً، لا مخمَّنةٌ: خمّنتُها أوّلَ
#    مرّةٍ فسقطت سبعُ نشراتٍ سليمةٍ على شاهدٍ لا وجودَ له.
NEED_EN = [("الرابط", "ahseyam.github.io/unified-lesson"),
           ("التوقيع", "PlanningandSchoolAccreditationDepartment"),
           ("Before you start", "Beforeyoustart"),
           ("Step by step", "Stepbystep"),
           ("Where things usually go wrong", "Wherethingsusuallygowrong"),
           ("Frequently asked", "Frequentlyasked"),
           ("شاهدُ اللقطة", "takenfromtheplatformitself")]
# ⛔ ولا يُقال «لا تسرّبَ» إلا بعد قياسٍ: هذه تُفحص في اللغتين
FORBID = [("لا كلمةَ سرٍّ", "1121986"), ("ولا بريدَ شخصيّ", "ahmed.mahmud")]


def check(docx, lang="ar"):
    pdf = docx[:-5] + ".pdf"
    name = os.path.basename(docx)
    res, ok = [], True

    txt = docx_text(docx).replace(" ", "")
    d = fitz.open(pdf)
    W, H = d[0].rect.width, d[0].rect.height

    def col(pg, r):
        px = pg.get_pixmap(clip=fitz.Rect(*r), dpi=60)
        s, n = px.samples, px.n
        idx = range(0, len(s) // n, 7)
        return round(sum(1 for j in idx
                         if max(s[j*n:j*n+3]) - min(s[j*n:j*n+3]) > 18
                         or max(s[j*n:j*n+3]) < 200) / len(idx) * 100)

    def say(label, good, extra=""):
        nonlocal ok
        ok &= bool(good)
        res.append("  %s %-34s %s" % ("✓" if good else "⛔", label, extra))

    say("القياس A4", abs(W/MM - 210) < 1 and abs(H/MM - 297) < 1,
        "%.0f×%.0f مم" % (W/MM, H/MM))
    say("عددُ الصفحات ٣ فأكثر", d.page_count >= 3, "%d صفحة" % d.page_count)

    heads = [col(pg, (15*MM, 3*MM, 195*MM, 22*MM)) for pg in d]
    foots = [col(pg, (15*MM, 280*MM, 195*MM, 294*MM)) for pg in d]
    say("الكليشةُ في كل صفحة", min(heads) > 20 and min(foots) > 20,
        "ترويسة %d–%d٪ · تذييل %d–%d٪" % (min(heads), max(heads), min(foots), max(foots)))

    over = 0
    for pg in d:
        for b in pg.get_text("blocks"):
            if not (b[4] or "").strip():
                continue
            if b[3] <= BAND:            # داخل شريط الكليشة
                continue
            if b[0] < SIDE - 2 or b[2] > W - SIDE + 2 or b[1] < BAND - 2 or b[3] > H - FOOT + 2:
                over += 1
    say("لا نصَّ يتجاوز الهوامش", over == 0, "%d كتلة" % over)

    shots = sum(len(pg.get_images()) for pg in d) - d.page_count   # عدا الكليشة
    say("اللقطاتُ مُدرَجة", shots >= 4, "%d صورة" % shots)

    # ── المحتوى من الـDOCX ──
    for label, v in (NEED_EN if lang == "en" else NEED_AR):
        say(label, v in txt)
    say("لا ذكرَ للبرنامج السابق", "الحصصالتطبيقية" not in txt)
    for label, v in FORBID:
        say(label, v not in txt)

    print("── %s ──" % name)
    print("\n".join(res))
    return ok


if __name__ == "__main__":
    good = True
    # ⚠️ اللغةُ من المسار: مجلَّدُ English نشراتُه إنجليزية
    for a in sys.argv[1:]:
        good &= check(a, "en" if os.sep + "English" + os.sep in a else "ar")
        print()
    raise SystemExit(0 if good else 1)
