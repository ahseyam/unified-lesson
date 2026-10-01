# -*- coding: utf-8 -*-
"""نشراتُ استخدام منصة الحصة الموحَّدة — خمسٌ، لكل دورٍ نشرتُه، بلقطاتٍ حيّة.

المحتوى في nshcontent.py واللقطاتُ في nshshots.py — وهذا الملفُّ يبني المستند
بهوية ابن خلدون: الكليشةُ خلفيةَ صفحةٍ، والعنوانُ في الشريط الأزرق، وخطُّ الجزيرة.

الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 nshbuild.py <المفتاح> <الملف.docx>
           والمفاتيح: teacher · peer · principal · deputy · supervisor
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image

from dox import (set_doc_defaults, set_compat15, rtl_par, par_space, par_keep,
                 par_shd, par_border_bottom, run, make_table, cell_text, cell_shd,
                 cell_valign, row_nosplit, normalize_tables, _add_pPr_el,
                 _insert_ordered, GENDER)
from band import band_title
import nshcontent as C

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, GREY, INK = "C00000", "7F7F7F", "1A1A1A"
NUMBG, SOFT, WARNBG, WARNLN = "E9EEF6", "F6F8FB", "FDF3F3", "E6B8B8"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
G = lambda m, f: f if GENDER == "f" else m
SHOTS = os.path.join(ROOT, "١٠ - أدلّة استخدام المنصة", "صور", G("بنين", "بنات"))

KEY = sys.argv[1]
OUT = sys.argv[2]
S = C.SHEETS[KEY]

doc = Document()
set_doc_defaults(doc)
set_compat15(doc)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.0)
sec.top_margin, sec.bottom_margin = Cm(2.85), Cm(2.7)
sec.header_distance, sec.footer_distance = Cm(0.2), Cm(1.6)
sec._sectPr.append(OxmlElement('w:bidi'))

# ── الكليشةُ خلفيةَ كل صفحة ──
hp = sec.header.paragraphs[0]
hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run()
r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=12, inner="C9D2DE", isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def para(sb=0, sa=3, keep=False):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, sb, sa)
    if keep:
        par_keep(p)
    return p


def heading(text, hint=None, sb=7):
    p = para(sb, 2, keep=True)
    par_border_bottom(p, TEAL, "8")
    run(p, "❖  " + text, size=11, bold=True, color=RED)
    if hint:
        run(p, "    " + hint, size=8, color=GREY, light=True)


# ⛔ **التوكيدُ `**نصّ**` كان يُطبع نجمتين.** كُتب في المحتوى بصيغة ماركداون
#    (سبعةُ مواضع)، و`run` يكتب النصَّ كما هو — فخرجت «نطاقُك **مجمعٌ كامل**»
#    في **ثلاثَ عشرةَ نشرةً منشورة** عربيةً وإنجليزية. فيُفسَّر هنا مرةً
#    واحدةً لكل كاتبٍ بدل حذفه من المحتوى. (١ أكتوبر ٢٠٢٦)
def runmd(p, text, size=9.6, color=INK, light=False):
    """يكتب النصَّ ويُثخّن ما بين نجمتين — ولا يطبع النجمتين."""
    for i, part in enumerate(str(text).split("**")):
        if part:
            run(p, part, size=size, color=color, bold=(i % 2 == 1),
                light=(light and i % 2 == 0))


def bullet(text, size=9.6, color=INK):
    p = para(0, 2)
    _add_pPr_el(p, 'ind', {'right': '220'})
    run(p, "•  ", size=8, color=TEAL)
    runmd(p, text, size=size, color=color)


ARN = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
arn = lambda n: str(n).translate(ARN)


def shot(name, max_w=W - 1.2):
    """⚠️ اللقطةُ تُدرج بعرضٍ ثابتٍ وارتفاعٍ محسوبٍ من نسبتها — وإلا شُوِّهت."""
    p = os.path.join(SHOTS, name + ".png")
    if not os.path.exists(p):
        raise SystemExit("⛔ لقطةٌ مفقودة: %s\n  شغّل: CLS_GENDER=%s python3 nshshots.py %s"
                         % (p, GENDER, name))
    iw, ih = Image.open(p).size
    h = max_w * ih / iw
    # ⚠️ سقفُ الارتفاع: لقطةٌ طويلةٌ تُزيح ما بعدها إلى صفحةٍ جديدةٍ وتترك فراغاً
    if h > 11.0:
        max_w, h = max_w * 11.0 / h, 11.0
    pp = para(2, 4)
    _add_pPr_el(pp, 'jc', {'val': 'center'})
    pp.add_run().add_picture(p, width=Cm(max_w), height=Cm(h))
    cap = para(0, 6)
    _add_pPr_el(cap, 'jc', {'val': 'center'})
    run(cap, "لقطةٌ من المنصة نفسِها", size=7.5, color=GREY, light=True)


def warn(text):
    t = make_table(doc, 1, [W], header_rows=0)
    frame(t, outer=WARNLN, osz=8)
    c = t.rows[0].cells[0]
    cell_shd(c, WARNBG)
    cell_valign(c, 'center')
    cell_text(c, "⛔  " + text, size=9.2, bold=True, color=RED)
    row_nosplit(t.rows[0])
    para(0, 4)


def box(title, lines, bg=TEALBG, line=TEAL):
    t = make_table(doc, 1 + len(lines), [W], header_rows=0)
    frame(t, outer=line, osz=8, inner=line, isz=2)
    cell_shd(t.rows[0].cells[0], line)
    cell_text(t.rows[0].cells[0], title, size=10, bold=True, color="FFFFFF")
    for r_ in t.rows[1:]:
        cell_shd(r_.cells[0], bg)
        row_nosplit(r_)
    for i, x in enumerate(lines):
        cell_text(t.rows[i + 1].cells[0], x, size=9.4, color=INK)
    para(0, 5)


# ═════════ الوجه الأول ═════════
band_title(sec, G("الحصة الموحَّدة", "الحصة الموحَّدة"))

p = para(0, 3)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED)
run(p, S["title"], size=13, bold=True, color=RED)
run(p, "  ◆", size=7, color=RED)

p = para(0, 6)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, S["who"], size=9.6, color=NAVY, bold=True)

p = para(2, 6)
runmd(p, S["why"], size=10, color=INK)

box(G("قبل أن تبدأ — ثلاثٌ تخصّ الجميع", "قبل أن تبدئي — ثلاثٌ تخصّ الجميع"),
    ["%s.  %s — %s" % (arn(i + 1), t, sub) for i, (t, sub, _d) in enumerate(C.COMMON_OPEN)])
for i, (t, _sub, d) in enumerate(C.COMMON_OPEN):
    bullet("%s %s" % (t + ":", d), size=9.2, color=GREY)

# ═════════ الخطوات ═════════
heading(G("خطوةً خطوة — بالصور", "خطوةً خطوة — بالصور"),
        "كلُّ لقطةٍ مأخوذةٌ من المنصة نفسِها، لا رسماً ولا تخيّلاً")

for i, (title, img, lines, wtext) in enumerate(S["steps"]):
    p = para(6 if i else 2, 2, keep=True)
    run(p, "%s" % arn(i + 1), size=13, bold=True, color=TEAL)
    run(p, "  " + title, size=11, bold=True, color=NAVY)
    for ln in lines:
        bullet(ln)
    if wtext:
        warn(wtext)
    if img:
        shot(img)

# ═════════ الوجه الأخير ═════════
heading("ما يقع فيه الخطأُ عادةً",
        "قبل أن تسأل — أكثرُ ما يتكرّر")
heads = ["ما يحدث", G("السبب وما تفعله", "السبب وما تفعلينه")]
t = make_table(doc, 1 + len(S["traps"]), [6.4, W - 6.4], header_rows=1)
frame(t)
for c in t.rows[0].cells:
    cell_shd(c, NAVY)
for j, x in enumerate(heads):
    cell_text(t.rows[0].cells[j], x, size=9.4, bold=True, color="FFFFFF")
for i, (a, b) in enumerate(S["traps"]):
    row_nosplit(t.rows[i + 1])
    cell_shd(t.rows[i + 1].cells[0], SOFT)
    cell_text(t.rows[i + 1].cells[0], a, size=9.2, bold=True, color=NAVY)
    cell_text(t.rows[i + 1].cells[1], b, size=9.2, color=INK)
para(0, 6)

heading(G("أسئلةٌ متكرّرة", "أسئلةٌ متكرّرة"))
for q, a in S["faq"]:
    p = para(2, 1)
    run(p, "س: ", size=9.4, bold=True, color=TEAL)
    run(p, q, size=9.4, bold=True, color=NAVY)
    p = para(0, 3)
    _add_pPr_el(p, 'ind', {'right': '260'})
    run(p, "ج: ", size=9.2, bold=True, color=TEAL)
    run(p, a, size=9.2, color=INK)

box(C.CLOSE_HELP[0], C.CLOSE_HELP[1], bg=SOFT, line=NAVY)

p = para(4, 0)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, C.SITE, size=9.6, bold=True, color=TEAL, latin=True)
p = para(0, 0)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "إدارة التخطيط والاعتماد المدرسي  ·  أ. أحمد صيام", size=8.4, color=GREY, light=True)

normalize_tables(doc)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("حُفظ: %s | %d خطوة | %d لقطة | %d مصيدة | %d سؤال"
      % (os.path.basename(OUT), len(S["steps"]),
         sum(1 for s in S["steps"] if s[1]), len(S["traps"]), len(S["faq"])))
