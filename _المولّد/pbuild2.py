# -*- coding: utf-8 -*-
"""نموذج تحضير الحصة — توأم استمارة الملاحظة الصفية في الإخراج (كليشة خلفيةً · العنوان في الشريط ·
إطار ٢ نقطة · صف كحلي لكل مجال). وجهان A4. أُعيد بناء المولّد 2026-09-27 مطابقاً للمنشور.

CLS_LESSON=<ملف فيه LESSON> ⇒ نموذج معبّأ استرشادي: تُكتب إجابات المعلم وتُؤشَّر خياراته.
الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f [CLS_LESSON=...] python3 pbuild2.py <الملف>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from dox import *
from dox import GENDER, _add_pPr_el, _trPr_set, _tcPr_set, _tblPr_set, _insert_ordered

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, HEADBG, NUMBG, ZEBRA, INNER = "C00000", "D9D9D9", "E9EEF6", "F6F8FB", "C9D2DE"
GREY, INK = "7F7F7F", "1A1A1A"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")

# ═════════ بيانات درس اختيارية ═════════
LESSON = {}
_lp = os.environ.get("CLS_LESSON", "")
if _lp:
    import importlib.util as _iu
    _sp = _iu.spec_from_file_location("lesson", _lp); _m = _iu.module_from_spec(_sp); _sp.loader.exec_module(_m)
    LESSON = _m.LESSON
    if os.environ.get("CLS_GENDER", "m") == "f":
        import lfem                      # ⚠️ مدارس البنات: كل فعل وضمير يعود على المعلمة أو المتعلمة يُؤنَّث
        LESSON = lfem.walk(LESSON)
ANS = "1F4E79"          # لون ما يكتبه المعلم
FIELDS = LESSON.get("fields", {})
TICKS = LESSON.get("ticks", {})

doc = Document(); set_doc_defaults(doc); set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.0)
s.top_margin, s.bottom_margin = Cm(2.85), Cm(2.7)
s.header_distance, s.footer_distance = Cm(0.2), Cm(1.6)
s._sectPr.append(OxmlElement('w:bidi'))

hp = s.header.paragraphs[0]; hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run(); r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=16, inner=INNER, isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def table(rows, widths, **kw):
    return frame(make_table(doc, rows, widths, header_rows=0), **kw)


def span(row, a, b, widths):
    c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
    _tcPr_set(c, 'tcW', {'w': str(int(sum(widths[a:b + 1]) / sum(widths) * 5000)), 'type': 'pct'})
    return c


LAB = 3.1
STAGE_H = 1.6 if LESSON else 2.55   # الفارغ يملأ الوجه الأول؛ والمعبّأ يتمدد بنصّه
LINE_ROOM = 1.6                    # سطر الخانة = المقاس × هذا؛ ‎1.32 كان يقصّ نوازل Sakkal
GRID = [LAB] + [(W - LAB) / 6] * 6


def gap(pt=3.5):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, pt)


def ticks(*items):
    return items


def tick_par(p, items, chosen, size, sep="     "):
    for i, it in enumerate(items):
        sel = it in chosen
        run(p, ("☑ " if sel else "☐ ") + it, size=size, bold=sel, color=ANS if sel else INK)
        if i < len(items) - 1:
            run(p, sep, size=size, color=INK)


def answer(label, idx=0, default=None):
    v = FIELDS.get(label, default)
    if isinstance(v, (list, tuple)):
        return v[idx] if idx < len(v) else None
    return v if idx == 0 else None


# ⛔ **عناوينُ المجالات تُقرأ من المصدر لا تُكتب بيد**: كان المطبوعُ يكتب
#    «المجالان ٢ و٣ · العرض وإدارة الحصة · مشاركة المتعلمين واندماجهم»
#    بينما `prepdef` يقول «المجال ٢ · العرض وإدارة الحصة» — فالورقُ يَعِد
#    بتغطية المجال الثالث **وله قسمٌ مستقلٌّ تحته**، وهو العنوانُ الذي قال
#    المصدرُ نفسُه إنه «يدّعي ما لا يفعل». و`printsrc` كان يحرس قوائمَ
#    الخيارات ولا يحرس العناوين. (قِيس ٢ أكتوبر ٢٠٢٦)
def _sec(key_n):
    """عنوانُ القسم ورقمُه من `prepdef` — بمطابقة وصفه، فلا يُكتب نصٌّ ثانٍ."""
    for x in _PD.SECTIONS:
        n = x.get("n") or ""
        if n.startswith(key_n):
            return x["t"], n
    raise SystemExit("⛔ لا قسمَ في prepdef وصفُه يبدأ بـ%r — "
                     "فلا يُكتب العنوانُ بيد." % key_n)


def domain_row(t, widths, num, name, h=0.46):
    r = t.rows[0]; row_height(r, h); row_nosplit(r)
    c = span(r, 0, len(widths) - 1, widths); c.text = ""
    pp = c.paragraphs[0]; rtl_par(pp); par_space(pp, 0, 0, 12)
    run(pp, f"{num}  ·  {name}", size=9.3, bold=True, color="FFFFFF")
    run(pp, "   — يقابله في الاستمارة " + num, size=7.5, color="DCE6F1", light=True)
    cell_shd(c, NAVY); cell_valign(c, 'center')


def domain_table(num, name, rows_spec):
    """rows_spec: [(التسمية, [(الامتداد, النص, نوع)], ارتفاع)] — النوع: blank · tick · note."""
    t = table(1 + len(rows_spec), GRID)
    if num: domain_row(t, GRID, num, name)
    for ri, (label, parts, h) in enumerate(rows_spec, 1):
        r = t.rows[ri]; row_height(r, h); row_nosplit(r)
        z = ZEBRA if ri % 2 == 0 else None
        cell_text(r.cells[0], label, size=8.2, bold=True, color=NAVY, shd=NUMBG, align='center')
        col = 1; fi = 0
        for sp, txt, kind in parts:
            cell = span(r, col, col + sp - 1, GRID)
            if kind == 'tick':
                p = cell_text(cell, "", size=8.2, color=INK, valign='center', shd=z,
                              line=round(8.2 * LINE_ROOM, 1))
                tick_par(p, txt, TICKS.get(label, ()), 8.2)
            else:
                va = 'center' if kind == 'note' else 'top'
                p = cell_text(cell, txt, size=7.5, color=GREY, valign=va, shd=z,
                              line=round(7.5 * LINE_ROOM, 1))
                a = answer(label, fi)      # الترتيب بين الخانات الكتابية وحدها
                if a:
                    run(p, ("  " if txt else "") + a, size=8, color=ANS)
                fi += 1
            col += sp
    if not num:
        t._tbl.remove(t.rows[0]._tr)
    return t


# ═════════ العنوان ═════════
from band import band_title
from prepdef import (TIMEMAP_K, TIME_KEYS, TIME_LEGACY, TIME_MAIN,
                     TIME_DIFF, TIME_DIFF_TITLE)
# ⛔ **قوائمُ المطبوع كانت تُكتب بيدها** فانحرفت عن المنصة: اثنتا عشرةَ قائمةً
#    منسوخةً، وأخطرُها «الإستراتيجية» — ستٌّ في الورق بينما المنصةُ تعرض بنكَ
#    التسعَ عشرةَ مقيَّداً بالاتجاه. و«استقصاء» و«مشروعات» **لا بطاقةَ لهما**،
#    و«تعلم تعاوني» تختلف رسماً عن «التعلم التعاوني» فلا تُطابَق آلياً — فتنقطع
#    سلسلةُ النظام على الورق: الإستراتيجيةُ المعلنةُ ← بطاقتُها ← درجةُ م٢·٣.
#    فصارت تُقرأ من `prepdef.TICKS`، وحارسٌ يمنع عودةَ النسخ. (١ أكتوبر ٢٠٢٦)
import prepdef as _PD


def _tk(key):
    """قائمةُ خياراتٍ من مصدرها الواحد — ولا تُكتب بيدٍ في مطبوع."""
    v = _PD.TICKS.get(key)
    if not v:
        raise SystemExit("⛔ لا قائمةَ باسم «%s» في prepdef.TICKS" % key)
    return tuple(v)
band_title(s, "نموذج تحضير الحصة")
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 3); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED)
run(p, LESSON.get("subtitle", "بمجالات استمارة الملاحظة الصفية وترتيبها"), size=9.5, bold=True, color=RED)
run(p, "  ◆", size=7, color=RED)

# ═════════ بيانات الحصة ═════════
DW = [1.9, 2.9, 1.7, 2.0, 2.25, 1.9, 1.85, 1.4, 2.2, 0.9]
t = table(2, DW)
info = LESSON.get("info", {})
for ri, labs in enumerate([("اسم المعلم", "المادة", "الصف/الشعبة", "التاريخ", "الحصة"),
                           ("موضوع الدرس", "رقم الدرس", "نسبة الإنجاز", "زمن الحصة", "عدد الطلاب")]):
    row_height(t.rows[ri], 0.5); row_nosplit(t.rows[ri])
    for j, lb in enumerate(labs):
        cell_text(t.rows[ri].cells[j * 2], lb, size=8, bold=True, color=NAVY, shd=NUMBG, align='center')
        unit = "٪" if lb == "نسبة الإنجاز" else ("د" if lb == "زمن الحصة" else "")
        c = t.rows[ri].cells[j * 2 + 1]
        if info.get(lb):
            p = cell_text(c, unit, size=8, color=GREY, align='end' if unit else 'center')
            run(p, ("  " if unit else "") + info[lb], size=8, bold=True, color=ANS)
        else:
            cell_text(c, unit, size=8, color=GREY, align='end')
gap(5)

# ═════════ المجال ١ ═════════
domain_table(*_sec("التخطيط وتجهيز بيئة التعلم"), rows_spec=[
    ("السؤال الأساسي", [(6, "", 'blank')], 0.5),
    ("أسئلة الوحدة", [(3, "١", 'blank'), (3, "٢", 'blank')], 0.5),
    ("أسئلة المحتوى", [(2, "١", 'blank'), (2, "٢", 'blank'), (2, "٣", 'blank')], 0.5),
    ("الهدف المعرفي", [(3, "١", 'blank'), (3, "٢", 'blank')], 0.5),
    ("الهدف المهاري", [(3, "١", 'blank'), (3, "٢", 'blank')], 0.5),
    ("الهدف الوجداني", [(3, "١", 'blank'), (3, "٢", 'blank')], 0.5),
    ("الاتجاه التدريسي", [(6, _tk("dir"), 'tick')], 0.44),
    ("عرض الأهداف", [(6, _tk("show"), 'tick')], 0.44),
    # ⛔ **لا قائمةَ ستّيّةٌ للإستراتيجية في الورق.** كانت ستّاً منسوخةً بيدها:
    #    «استقصاء» و«مشروعات» بلا بطاقةٍ في البنك أصلاً، و«تعلم تعاوني» تختلف
    #    رسماً عن «التعلم التعاوني» فلا تُطابَق. فتنقطع سلسلةُ النظام على الورق:
    #    من أعلنها لا تُشخَّص بطاقتُه ولا تُشتقّ درجةُ م٢·٣. والبنكُ تسعَ عشرةَ
    #    بطاقةً مقيَّدةً بالاتجاه — لا تُطبع خاناتٍ، فيُكتب اسمُها. (١ أكتوبر ٢٠٢٦)
    ("الإستراتيجية", [(6, "من بنك البطاقات التسع عشرة — ومقيَّدةٌ باتجاهك أعلاه", 'blank')], 0.44),
    ("مصادر التعلم", [(6, _tk("src"), 'tick')], 0.44),
    ("داعمات البيئة", [(6, _tk("env"), 'tick')], 0.44),
    ("أوراق العمل", [(6, _tk("sheets"), 'tick')], 0.44),
])
gap()

# ═════════ خريطة الزمن ═════════
# ⛔ جدولان لا جدولٌ واحد: لو صُفَّت خانتا التمايز مع صف المجموع قرأها القارئُ
#    جامعةً فبلغ ٦١ دقيقةً والحصةُ ٤٥. فهما في مجموعةٍ مستقلةٍ تحت عنوانٍ يقول
#    إنهما من زمن التنفيذ لا تُضافان إليه.
_LBL = dict(TIMEMAP_K)
_tm = LESSON.get("time", [])
_tmv = {TIME_LEGACY[j]: v for j, v in enumerate(_tm) if j < len(TIME_LEGACY)}
# ⛔ CLS_MARK: عناوينُ الخلايا التي عُدِّلت، تُظلَّل صفراء في نسخة المراجعة وحدها.
_MARK = [x.strip() for x in os.environ.get("CLS_MARK", "").split("،") if x.strip()]
MARKBG = "FFF2A8"


def _timetab(keys, widths, care=False):
    tb = table(2, widths)
    for c, kk in enumerate(keys):
        lb = _LBL[kk]
        hit = any(m in lb for m in _MARK)
        cell_text(tb.rows[0].cells[c], lb, size=7.8, bold=True,
                  color=TEAL if care else NAVY, align='center',
                  shd=MARKBG if hit else (TEALBG if care else HEADBG))
        pp = cell_text(tb.rows[1].cells[c], "د", size=8, color=GREY, align='end',
                       shd=MARKBG if hit else (TEALBG if care else None))
        v = _tmv.get(kk, "")
        if v:
            run(pp, "  " + str(v), size=8.5, bold=True, color=ANS)
    row_height(tb.rows[0], 0.44); row_nosplit(tb.rows[0])
    row_height(tb.rows[1], 0.5); row_nosplit(tb.rows[1])
    return tb


_timetab(TIME_MAIN, [W / 5] * 5)
gap(3)
_ht = doc.add_paragraph()
run(_ht, TIME_DIFF_TITLE + " — داخلَه لا يُضافان إليه:", size=7.6, bold=True, color=TEAL)
par_space(_ht, after=1)
_timetab(TIME_DIFF, [W / 5, W / 5], care=True)
gap(5)

# ═════════ المجال ٢ ═════════
SW = [2.1, 6.2, 6.2, 4.5]
t = table(6, SW)
domain_row(t, SW, *_sec("العرض وإدارة الحصة"))
for i, lb in enumerate(["المرحلة", "ما يفعله المعلم", "ما يفعله المتعلم — فعلٌ يُرى", "نمط العمل وتقويمه"]):
    cell_text(t.rows[1].cells[i], lb, size=8.2, bold=True, color=TEAL if i == 2 else NAVY, align='center',
              shd=TEALBG if i == 2 else HEADBG)
row_height(t.rows[1], 0.46); row_nosplit(t.rows[1])
for ri, st in enumerate(["التهيئة\nمدخل الحصة", "النشاط الأول", "النشاط الثاني", "الغلق\nعلم بالقلم"], 2):
    r = t.rows[ri]; row_height(r, STAGE_H); row_nosplit(r)
    stage = LESSON.get("stages", {}).get(st.split("\n")[0], {})
    cell_lines(r.cells[0], [(x, True) for x in st.split("\n")] +
               [("الزمن: " + stage.get("زمن", "....") + " د", False)],
               size=8.2, color=NAVY, shd=NUMBG, valign='center', align='center', space=1)
    for ci, key in ((1, "المعلم"), (2, "المتعلم")):
        p = cell_text(r.cells[ci], "", size=5, valign='top', shd="F3F9FA" if ci == 2 else None)
        if stage.get(key):
            par_space(p, 0, 0, round(8 * LINE_ROOM, 1))
            run(p, stage[key], size=8, color=ANS)
    ch = stage.get("نمط", ())
    c3 = r.cells[3]; c3.text = ""
    for li, items in enumerate((("فردي", "ثنائي", "جماعي"), ("تقويم ذاتي", "تقويم أقران"),
                                ("تقويم بنائي", "علاجي"))):
        p = c3.paragraphs[0] if li == 0 else c3.add_paragraph()
        rtl_par(p); _add_pPr_el(p, 'jc', {'val': 'center'})
        par_space(p, 0, 5 if li < 2 else 0, round(7 * LINE_ROOM, 1))
        tick_par(p, items, ch, 7, sep="   ")
    cell_valign(c3, 'center')

# ═════════ الوجه الثاني: يبدأ بجدول المهام ═════════
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 2); par_break_before(p)
domain_table(*_sec("تابع: التمايز والتحفيز والربط بالحياة"), rows_spec=[
    ("المهمة المكيَّفة", [(6, "للفئات الأولى بالرعاية — ما هي، ولمن، ومتى تُسلَّم:", 'note')], 0.95),
    ("المهمة الإثرائية", [(6, "لسريعي التعلم والموهوبين — تنمّي تفكيراً أعلى لا مزيداً من التمارين:", 'note')], 0.95),
    ("التحفيز والمشاركة", [(6, _tk("motiv"), 'tick')], 0.95),
    ("الربط بالحياة", [(6, "مثال أو تطبيق من حياة الطلاب:", 'note')], 0.95),
])
gap(5)
# ⛔ **المجالُ ٣ كان بلا خانةٍ في الورق كما كان في المنصة** — ستةُ مؤشراتٍ
#    (٢٤ درجة) تُرصد على الطالب بلا تخطيطٍ لها. (١ أكتوبر ٢٠٢٦)
# ⚠️ **وتسمياتُه تُقرأ من `prepdef` لا تُكتب هنا**: كتبتُها بيدي أولَ مرةٍ
#    فخرجت نسخةُ البنات بعربيّةٍ مكسورة («مشاركة المتعلمات واندماجهم» ·
#    «ما سيقوله المتعلمة ويفعله») — لأن المؤنِّثَ العامَّ لا يبلغ ما تبلغه
#    صيغتان مكتوبتان في المصدر. وهي المخالفةُ نفسُها التي جئتُ أُصلحها.
_s3 = next(x for x in _PD.SECTIONS if x["t"] == "المجال ٣")
domain_table(_s3["t"], _s3["n"],
             [(r["label"], [(6, r.get("note", "") + ":", 'note')], 0.95) for r in _s3["rows"]])
gap(5)
domain_table(*_sec("المهارات المستهدفة"), rows_spec=[
    ("القراءة والكتابة", [(3, ticks("نص يقرؤه الطلاب", "كتابة إجابة أو جملة"), 'tick'), (3, "كيف:", 'note')], 0.65),
    ("المهارات العددية", [(3, ticks("حساب ينفّذه الطلاب", "قراءة بيانات"), 'tick'), (3, "كيف:", 'note')], 0.65),
    ("مهارات التفكير", [(6, _tk("thinkk"), 'tick')], 0.44),
    ("نشاط التفكير", [(6, "السؤال أو المهمة التي تنمّي التفكير:", 'note')], 0.9),
])
gap(5)
domain_table(*_sec("التقنية والتعلم الإلكتروني"), rows_spec=[
    ("أدوات التقييم", [(6, _tk("tools"), 'tick')], 0.46),
    ("توظيف التقنية", [(6, _tk("tech"), 'tick')], 0.44),
    ("المهمة على المنصة", [(6, "اسم المنصة والمهمة المُسنَدة:", 'note')], 0.75),
])
gap(5)
domain_table(*_sec("التقويم والتغذية الراجعة"), rows_spec=[
    ("التقويم التشخيصي", [(6, "سؤال أو نشاط قبلي يكشف المعرفة السابقة:", 'note')], 0.8),
    ("التقويم البنائي", [(6, "بعد كل هدف — الأداة:", 'note')], 0.8),
    ("الذاتي والأقران", [(3, _tk("selfpeer"), 'tick'), (3, "الأداة أو المعيار:", 'note')], 0.7),
    ("التغذية الراجعة", [(6, "اللحظة التي يوظّف فيها الطالب الملاحظة لتحسين عمله:", 'note')], 0.8),
    ("التقييم الختامي", [(3, "س١", 'blank'), (3, "س٢", 'blank')], 0.95),
    ("أسئلة تفكير عليا", [(3, "س٣", 'blank'), (3, "س٤", 'blank')], 0.95),
])
gap(5)
domain_table(*_sec("بناء شخصية الطالب"), rows_spec=[
    ("قيمة الأسبوع", [(6, "وربطها بمحتوى الدرس:", 'note')], 0.75),
    ("مهارة التلخيص", [(6, _tk("summar"), 'tick')], 0.44),
    ("مهارة الغلق", [(6, _tk("close"), 'tick')], 0.44),
])

p = doc.add_paragraph(); rtl_par(p); par_space(p, 2, 3, 10)
run(p, "المجال ٨ (مهارات المعلم وقدراته) يُرصد في الزيارة ولا يُحضَّر.", size=7.5, color=GREY, light=True)
if LESSON.get("source"):
    run(p, "     مصدر الدرس: " + LESSON["source"], size=7.5, color=GREY, light=True)

# ═════════ التوقيعات ═════════
GW = [W / 4] * 4
t = table(2, GW)
for i, lb in enumerate(["توقيع المعلم", "الوكيل التعليمي", "المشرف المختص", "مدير المدرسة"]):
    cell_text(t.rows[0].cells[i], lb, size=8, bold=True, color=NAVY, shd=NUMBG, align='center')
    cell_text(t.rows[1].cells[i], "", size=9)
row_height(t.rows[0], 0.44); row_nosplit(t.rows[0])
row_height(t.rows[1], 1.0); row_nosplit(t.rows[1])

# فقرة ختامية بأصغر ارتفاع: وورد يبقي فقرة بعد آخر جدول
_p = doc.add_paragraph(); rtl_par(_p); par_space(_p, 0, 0, 1); run(_p, "", size=1)

# ═════════ التذييل ═════════
fp = s.footer.paragraphs[0]; fp.text = ""
rtl_par(fp); _add_pPr_el(fp, 'jc', {'val': 'center'}); par_space(fp, 0, 0)
run(fp, "نموذج تحضير الحصة — مدارس ابن خلدون   ·   الوجه ", size=7, color=GREY, light=True)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), ' PAGE ')
fr = OxmlElement('w:r'); frpr = OxmlElement('w:rPr'); rf = OxmlElement('w:rFonts')
for a in ('ascii', 'hAnsi', 'cs'): rf.set(qn('w:' + a), FONT)
frpr.append(rf)
for tg, v in (('sz', '14'), ('szCs', '14'), ('color', GREY)):
    e = OxmlElement('w:' + tg); e.set(qn('w:val'), v); frpr.append(e)
fr.append(frpr); tt = OxmlElement('w:t'); tt.text = "1"; fr.append(tt)
fld.append(fr); fp._p.append(fld)

normalize_tables(doc)
out = sys.argv[1] if len(sys.argv) > 1 else "p2.docx"
doc.save(out); print("حُفظ:", out)
