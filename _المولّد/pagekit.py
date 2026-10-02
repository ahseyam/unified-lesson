# -*- coding: utf-8 -*-
"""إطارُ الصفحة الطولية بهوية ابن خلدون — **مصدرٌ واحدٌ لكل بانٍ**.

⛔ كانت كتلةُ التهيئة (مقاسُ A4 · الهوامش · الكليشةُ خلفيةَ كل صفحة بمرساة
   `behindDoc`) منسوخةً في كل بانٍ. والمنسوخُ يفترق: يُضبط الهامشُ في واحدٍ
   ولا يُضبط في الآخر، فتخرج مطبوعاتُ مشروعٍ واحدٍ بهويّتين.
   (وقاعدةُ الذاكرة: استعِر دوالّ DOCX ولا تكتبها من جديد.)

الاستعمال:
    from pagekit import portrait
    doc = Document(); portrait(doc)
"""
import os

from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

from dox import par_space, set_compat15, set_doc_defaults

HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")


def portrait(doc, klisha=None, top=2.85, bottom=2.7, side=1.0):
    """A4 طوليّ · هوامشُ الكليشة · والكليشةُ خلفيةَ كل صفحةٍ مقفولةً.

    ⚠️ الصورةُ تُحوَّل من `inline` إلى `anchor` بـ`behindDoc="1"` — وبدونها
       تُزيح النصَّ بدل أن تقع خلفه. و`locked` يمنع جرَّها بالخطأ في وورد.
    """
    kl = klisha or KLISHA
    if not os.path.exists(kl):
        raise SystemExit("⛔ لا كليشة: %s" % kl)
    set_doc_defaults(doc)
    set_compat15(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(side)
    sec.top_margin, sec.bottom_margin = Cm(top), Cm(bottom)
    sec.header_distance, sec.footer_distance = Cm(0.2), Cm(1.6)
    sec._sectPr.append(OxmlElement('w:bidi'))

    hp = sec.header.paragraphs[0]
    hp.text = ""
    par_space(hp, 0, 0)
    r = hp.add_run()
    r.add_picture(kl, width=Cm(21.0), height=Cm(29.7))
    inline = r._r.find('.//' + qn('wp:inline'))
    graphic = inline.find(qn('a:graphic'))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" '
        'simplePos="0" relativeHeight="0" behindDoc="1" locked="1" '
        'layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="7560000" cy="10692000"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/>'
        '</wp:anchor>')
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)
    return sec
