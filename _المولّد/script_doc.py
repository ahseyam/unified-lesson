# -*- coding: utf-8 -*-
"""سيناريو التسجيل — نصُّ القراءة مطابقٌ لترقيم الشرائح.

يُبنى من `deckcontent.SLIDES` نفسه، فلا يفترق النصّ عن الشريحة.
لكل شريحة: رقمها وعنوانها · ما يظهر على الشاشة · النصّ المقروء بخطٍّ كبير مريح · زمنها.
الاستعمال: CLS_FONTSET=js python3 script_doc.py <الملف.docx>
"""
import os
import sys

from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dox import *                                   # noqa: F401,F403
from dox import _add_pPr_el
from band import band_title
from deckcontent import SLIDES

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, HEADBG, GREY, INK = "C00000", "E9EEF6", "7F7F7F", "1A1A1A"
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
AR = "٠١٢٣٤٥٦٧٨٩"


def ar(n):
    return "".join(AR[int(d)] for d in str(n))


doc = Document()
set_doc_defaults(doc)
set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(3.3), Cm(1.6)
s.left_margin = s.right_margin = Cm(1.5)
W = 18.0

# ═════════ الكليشة خلفيةً لكل صفحة (كما في بقية المطبوعات) ═════════
hp = s.header.paragraphs[0]
rtl_par(hp)
par_space(hp, 0, 0, 1)
r = hp.add_run()
r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)
band_title(s, "سيناريو التسجيل")


def gap(pt=6):
    p = doc.add_paragraph()
    rtl_par(p)
    par_space(p, 0, 0, pt)


def head(t, sub=None):
    p = doc.add_paragraph()
    rtl_par(p)
    par_space(p, 4, 3)
    _add_pPr_el(p, "jc", {"val": "center"})
    run(p, t, size=13, bold=True, color=NAVY)
    if sub:
        q = doc.add_paragraph()
        rtl_par(q)
        par_space(q, 0, 6)
        _add_pPr_el(q, "jc", {"val": "center"})
        run(q, sub, size=9.5, color=TEAL)


total = sum(x["d"] for x in SLIDES)
head("نظام الحصة الموحَّدة — سيناريو التسجيل",
     f"{ar(len(SLIDES))} شريحة  ·  زمن القراءة التقديري {ar(total // 60)} دقيقة و{ar(total % 60)} ثانية")

# ═════════ إرشادات التسجيل ═════════
t = make_table(doc, 2, [W], borders_color="C9D2DE")
cell_text(t.rows[0].cells[0], "قبل أن تبدأ", size=10, bold=True, color="FFFFFF", shd=NAVY, align='center')
cell_lines(t.rows[0].cells[0] if False else t.rows[1].cells[0], [
    ("افتح ملف العرض بالنقر المزدوج، ثم اضغط F لملء الشاشة — والأسهم أو المسافة للانتقال.", False),
    ("كل فقرةٍ هنا تقابل شريحةً برقمها، فانتقل إلى الشريحة التالية عند نهاية فقرتها.", False),
    ("الأزمنة تقديرية بقراءةٍ متأنّية؛ وما بين القوسين إرشادٌ لك لا يُقرأ.", False),
    ("إن أخطأت فلا تُعد التسجيل من أوله: قف، واصمت ثانيتين، وأعد الجملة — يسهل قصّها بعدُ.", False),
], size=9.5, color=INK, space=3)
row_height(t.rows[0], 0.55)
gap(8)

# ═════════ الشرائح ═════════
for i, sl in enumerate(SLIDES, 1):
    rows = make_table(doc, 3, [2.6, W - 2.6], borders_color="C9D2DE")
    # الرأس: الرقم والعنوان والزمن
    cell_text(rows.rows[0].cells[0], "الشريحة " + ar(i), size=10, bold=True,
              color="FFFFFF", shd=NAVY, align='center')
    p = cell_text(rows.rows[0].cells[1], sl["t"], size=10.5, bold=True, color=NAVY, shd=HEADBG)
    run(p, "   ·   " + sl["s"], size=9, color=TEAL)
    run(p, "   ·   نحو " + ar(sl["d"]) + " ثانية", size=8.5, color=GREY, light=True)
    row_height(rows.rows[0], 0.5)
    row_nosplit(rows.rows[0])
    # ما يظهر على الشاشة
    cell_text(rows.rows[1].cells[0], "على الشاشة", size=8.5, bold=True, color=TEAL,
              shd=TEALBG, align='center')
    scr = [("◆  " + x, False) for x in (sl.get("b") or [])] + \
          [("◆  " + x, False) for x in (sl.get("x") or [])]
    if sl.get("cap"):
        scr.append(("🖼  الصورة: " + sl["cap"], True))      # ما يراه المشاهد في الصورة
    cell_lines(rows.rows[1].cells[1], scr, size=8.5, color=GREY, space=2)
    row_nosplit(rows.rows[1])
    # النصّ المقروء
    cell_text(rows.rows[2].cells[0], "اقرأ", size=10, bold=True, color="FFFFFF",
              shd=RED, align='center')
    q = cell_text(rows.rows[2].cells[1], "", size=6)
    par_space(q, 2, 2, 22)
    run(q, sl["n"], size=13, color=INK)
    row_nosplit(rows.rows[2])
    gap(7)

p = doc.add_paragraph()
rtl_par(p)
par_space(p, 4, 0)
_add_pPr_el(p, "jc", {"val": "center"})
run(p, "انتهى السيناريو  ·  مدارس ابن خلدون", size=9, color=GREY, light=True)

normalize_tables(doc)
out = sys.argv[1] if len(sys.argv) > 1 else "script.docx"
doc.save(out)
print("حُفظ:", os.path.basename(out), "|", len(SLIDES), "فقرة |",
      f"{ar(total // 60)}:{ar(total % 60):0>2}", "دقيقة")
