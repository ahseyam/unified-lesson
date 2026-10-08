# -*- coding: utf-8 -*-
"""تقريرُ تسجيل الحصص الموحَّدة لأسبوعٍ بعينه — لكل مدرسةٍ ولكل حصة.

⛔ **ولماذا؟** يُصدِره مديرُ المنصة يومَ الثلاثاء السابقَ لأسبوع التنفيذ، فتثبت
   به الحصصُ وتُبنى عليه خطةُ تحرّك المقيّمين الخارجيين بين المدارس والمجمعات
   (قرارُ المستشار ٨ أكتوبر ٢٠٢٦، وأعلنه للوكلاء: «الساعة ١:٠٠ ظهراً يصدر
   تقريرٌ خاصٌّ باكتمال تسجيل المدارس لحصص الأسبوع الثامن»).

⚠️ و«مكتملةٌ» حدُّها **عقدُ الحجز** نفسُه الذي في المنصة: المعلم · الفصل ·
   الاتجاه · الإستراتيجية. ولا يُخترع هنا حدٌّ ثانٍ، وإلّا اختلف الورقُ عن
   الشاشة فسقطت الثقةُ بهما معاً.
⚠️ و«٠» و«-» و«.» ليست قيماً — القاعدةُ عينُها في `realVal` بالصفحة.
⚠️ ويُبنى من **المخزن الحيّ** لا من نسخةٍ على القرص: تقريرٌ بُني على أمسِ
   يُصدَر اليومَ فيُحاسَب به من أتمَّ عملَه ليلاً.

الاستعمال:
    python3 weekrep.py "الأسبوع الثامن" [مجلَّدُ المخرَج]
"""
import datetime
import json
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from dox import (run, rtl_par, par_space, par_shd, make_table, cell_text, cell_shd,
                 row_height, row_nosplit, normalize_tables, par_border_bottom)
import pagekit

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRV = os.path.expanduser("~/Desktop/خادم الحصة الموحَّدة (خاصّ — لا يُنشر)/store.json")
NAVY, TEAL, RED, OKC = "355E91", "2F7F95", "C00000", "2E7D32"
HEADBG, ZEBRA, WARNBG = "E9EEF6", "F6F8FB", "FFF4E0"

NOTVAL = {"0", "٠", "-", "—", "–", "_", ".", "،", "/", "×", "لا", "na", "n/a"}
NEED = [("teacher", "المعلم"), ("klass", "الفصل"),
        ("approach", "الاتجاه"), ("strategy", "الإستراتيجية")]


def val(v):
    t = re.sub(r"[\s‏‎]+", "", str(v or ""))
    return bool(t) and t.lower() not in NOTVAL


def gaps(x):
    return [a for f, a in NEED if not val(x.get(f))]


def arn(n):
    return str(n).translate(str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩"))


def pull(sid):
    cfg = json.load(open(SRV, encoding="utf-8"))
    r = subprocess.run(["curl", "-sS", "--max-time", "300",
                        "%s?kind=platform&id=%s&key=%s" % (cfg["url"], sid, cfg["join"])],
                       capture_output=True, text=True)
    d = json.loads(r.stdout)
    if not d.get("ok"):
        raise SystemExit("⛔ تعذّرت القراءة: %s" % str(d.get("error"))[:120])
    return d["data"]


def build(sid, gender, week, out):
    os.environ["CLS_GENDER"] = gender
    data = pull(sid)
    sch = [x for x in (data.get("sched") or []) if x.get("week") == week]
    prep = data.get("prep") or {}
    doc = Document()
    pagekit.portrait(doc)

    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 2)
    run(p, "تقريرُ تسجيل الحصص الموحَّدة", size=19, bold=True, color=NAVY)
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 10)
    run(p, week + " · " + ("قسم البنين" if gender == "m" else "قسم البنات"),
        size=13, bold=True, color=TEAL)
    par_border_bottom(p)

    full = [x for x in sch if not gaps(x)]
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 6, 8)
    run(p, "المسجَّل: ", size=12, bold=True)
    run(p, arn(len(sch)) + " حصة", size=12, bold=True, color=NAVY)
    run(p, "   ·   مكتملةُ البيانات: ", size=12, bold=True)
    run(p, arn(len(full)), size=12, bold=True, color=OKC)
    run(p, "   ·   ناقصة: ", size=12, bold=True)
    run(p, arn(len(sch) - len(full)), size=12, bold=True, color=RED)
    if sch:
        run(p, "   ·   نسبةُ الاكتمال: " + arn(round(100 * len(full) / len(sch))) + "٪",
            size=12, bold=True, color=TEAL)

    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 8)
    run(p, "والحصةُ «مكتملةٌ» إذا تمَّت أربعتُها: اسمُ المعلم · الفصل · "
           "الاتجاه التدريسي · الإستراتيجية. وما نقص منها لا تُحجز به الخانة.",
        size=9.5, color="666666")

    # ── جدولُ المدارس ──
    by = {}
    for x in sch:
        by.setdefault((x.get("sector") or "—", x.get("complex") or "—",
                       x.get("stage") or "—"), []).append(x)
    rows = [["القطاع", "المجمع", "المدرسة", "مسجَّلة", "مكتملة", "ناقصة", "الاكتمال"]]
    for k in sorted(by):
        g = by[k]
        f = [x for x in g if not gaps(x)]
        rows.append([k[0], k[1], k[2], arn(len(g)), arn(len(f)),
                     arn(len(g) - len(f)), arn(round(100 * len(f) / len(g))) + "٪"])
    # ⚠️ `make_table` يأخذ **عددَ الصفوف** لا بياناتِها — والخاناتُ تُملأ بعدها
    t = make_table(doc, len(rows), [1.9, 2.1, 5.0, 1.9, 1.9, 1.7, 2.0])
    for i, r0 in enumerate(t.rows):
        row_height(r0, 0.75 if i else 0.85)
        row_nosplit(r0)
        for j, c in enumerate(r0.cells):
            cell_text(c, rows[i][j], size=10 if i else 10.5, bold=(i == 0),
                      color=NAVY if i == 0 else None, align="center")
            if i == 0: cell_shd(c, HEADBG)
            elif i % 2 == 0: cell_shd(c, ZEBRA)

    # ── تفصيلُ كل مدرسة ──
    for k in sorted(by):
        g = sorted(by[k], key=lambda y: (y.get("day") or "", y.get("period") or ""))
        p = doc.add_paragraph(); rtl_par(p); par_space(p, 12, 4)
        run(p, k[2] + " — " + k[1] + " · " + k[0], size=13, bold=True, color=NAVY)
        rr = [["اليوم", "الحصة", "المعلم", "المادة", "الحالة", "الناقص"]]
        for x in g:
            gp = gaps(x)
            rr.append([x.get("day") or "—", x.get("period") or "—",
                       (x.get("teacher") or "—").strip(),
                       (x.get("subject") or x.get("spec") or "—").strip(),
                       "مكتملة" if not gp else "ناقصة",
                       " · ".join(gp) if gp else "—"])
        t2 = make_table(doc, len(rr), [2.4, 1.9, 4.3, 3.2, 1.9, 3.1])
        for i, r0 in enumerate(t2.rows):
            row_height(r0, 0.7 if i else 0.8)
            row_nosplit(r0)
            for j, c in enumerate(r0.cells):
                cell_text(c, rr[i][j], size=9.5 if i else 10,
                          bold=(i == 0 or (i and j == 4 and rr[i][4] == "ناقصة")),
                          color=NAVY if i == 0 else (RED if (i and j == 4 and rr[i][4] == "ناقصة") else None),
                          align="center")
                if i == 0: cell_shd(c, HEADBG)
                elif rr[i][4] == "ناقصة": cell_shd(c, WARNBG)

    p = doc.add_paragraph(); rtl_par(p); par_space(p, 14, 0)
    run(p, "صدر في " + datetime.datetime.now().strftime("%Y-%m-%d") + " الساعة "
           + datetime.datetime.now().strftime("%H:%M") + " — من المخزن المشترك مباشرةً.",
        size=9.5, color="666666")

    normalize_tables(doc)
    doc.save(out)
    return len(sch), len(full)


def main():
    week = sys.argv[1] if len(sys.argv) > 1 else "الأسبوع الثامن"
    # ⛔ **ولا يُكتب التقريرُ في شجرة التسليم**: فيها أسماءُ معلمين وحصصُهم،
    #    والمستودعُ عامٌّ منشور. فمخرَجُه خارجَها على سطح المكتب.
    outdir = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser(
        "~/Desktop/تقارير الحصة الموحَّدة")
    os.makedirs(outdir, exist_ok=True)
    tot = []
    for sid, g, nm in (("ikm_db", "m", "بنين"), ("ikf_db", "f", "بنات")):
        out = os.path.join(outdir, "تقرير تسجيل الحصص — %s — %s.docx" % (week, nm))
        n, f = build(sid, g, week, out)
        tot.append((nm, n, f, out))
        print("  ✓ %s: %d حصة · مكتملة %d (%d%%) → %s"
              % (nm, n, f, round(100 * f / max(1, n)), os.path.basename(out)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
