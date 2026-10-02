# -*- coding: utf-8 -*-
"""سيناريو تسجيل فيديو رحلة الدور — مستندٌ لكل دور، بهوية ابن خلدون.

⛔ المشاهدُ تُقرأ من `nshcontent` (خطواتُ النشرة نفسُها) — فلا يَعِد الفيديو
   بما لا يقوله الدليل، ولا يُكتب النصُّ مرّتين. وما يخصُّ التسجيلَ وحدَه
   في `vidcontent.py`.

⚠️ ولقطةُ كل مشهدٍ هي لقطةُ النشرة نفسُها، فيعرف المسجِّلُ الشاشةَ المقصودة
   قبل أن يفتحها — ولا تُصنع لقطةٌ ثانيةٌ تفترق عنها.

الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 vidbuild.py <المفتاح> <الملف.docx>
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from PIL import Image

from band import band_title
from dox import (GENDER, _add_pPr_el, cell_shd, cell_text, cell_valign,
                 make_table, normalize_tables, par_border_bottom, par_keep,
                 par_shd, par_space, row_nosplit, rtl_par, run)
from pagekit import portrait
import nshcontent as C
import vidcontent as V

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, GREY, INK = "C00000", "7F7F7F", "1A1A1A"
NUMBG, SOFT, WARNBG, WARNLN = "E9EEF6", "F6F8FB", "FDF3F3", "E6B8B8"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
G = lambda m, f: f if GENDER == "f" else m
SHOTS = os.path.join(ROOT, "١٠ - أدلّة استخدام المنصة", "صور", G("بنين", "بنات"))

KEY, OUT = sys.argv[1], sys.argv[2]
S = C.SHEETS[KEY]
SETUP, OPEN_SAY, CLOSE_SAY = V.ROLE[KEY]

doc = Document()
SEC = portrait(doc)

ARN = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
arn = lambda n: str(n).translate(ARN)


def mmss(sec):
    return arn("%d:%02d" % (sec // 60, sec % 60))


def para(sb=0, sa=3, keep=False):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, sb, sa)
    if keep:
        par_keep(p)
    return p


def runmd(p, text, size=9.6, color=INK, light=False):
    for i, part in enumerate(str(text).split("**")):
        if part:
            run(p, part, size=size, color=color, bold=(i % 2 == 1),
                light=(light and i % 2 == 0))


def heading(text, hint=None, sb=7):
    p = para(sb, 2, keep=True)
    par_border_bottom(p, TEAL, "8")
    run(p, "❖  " + text, size=11, bold=True, color=RED)
    if hint:
        run(p, "    " + hint, size=8, color=GREY, light=True)


def bullet(text, size=9.6, color=INK):
    p = para(0, 2)
    _add_pPr_el(p, 'ind', {'right': '220'})
    run(p, "•  ", size=8, color=TEAL)
    runmd(p, text, size=size, color=color)


def box(title, lines, bg=SOFT, line=NAVY):
    t = make_table(doc, 1 + len(lines), [W], borders_color=line, header_rows=0)
    cell_shd(t.rows[0].cells[0], bg)
    cell_text(t.rows[0].cells[0], title, size=9.6, bold=True, color=NAVY)
    for i, x in enumerate(lines, 1):
        cell_shd(t.rows[i].cells[0], bg)
        c = t.rows[i].cells[0]
        c.text = ""
        p = c.paragraphs[0]; rtl_par(p); par_space(p, 1, 1)
        run(p, "•  ", size=8, color=TEAL)
        runmd(p, x, size=9.2, color=INK)
        cell_valign(c, 'center')
        row_nosplit(t.rows[i])
    para(0, 5)
    return t


def shot(name, max_w=W - 5.0):
    p = os.path.join(SHOTS, name + ".png")
    if not os.path.exists(p):
        raise SystemExit("⛔ لقطةٌ مفقودة: %s" % p)
    iw, ih = Image.open(p).size
    h = max_w * ih / iw
    if h > 8.0:
        max_w, h = max_w * 8.0 / h, 8.0
    pp = para(1, 2)
    _add_pPr_el(pp, 'jc', {'val': 'center'})
    pp.add_run().add_picture(p, width=Cm(max_w), height=Cm(h))


# ═════════ العنوان ═════════
band_title(SEC, G("سيناريو تسجيل فيديو", "سيناريو تسجيل فيديو"))
p = para(2, 1)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, S["title"].replace("دليل استخدام المنصة — ", "رحلة: "), size=13, bold=True, color=NAVY)

TOTAL = V.T_OPEN + V.T_STEP * len(S["steps"]) + V.T_CLOSE
p = para(0, 6)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "الزمنُ المستهدف: %s دقيقة  ·  %s مشهداً  ·  %s" % (
    mmss(TOTAL), arn(len(S["steps"]) + 2), S["who"].replace("لمن: ", "")),
    size=9.4, color=TEAL, bold=True)

# ═════════ قبل التسجيل ═════════
heading("قبل أن تضغط «تسجيل»", "خمسُ قواعدَ لكل فيديو، وما يخصُّ هذا الدور")
box("قواعدُ التسجيل", ["**%s** (%s) — %s" % (t, s, d) for t, s, d in V.BEFORE])
box(G("وما يُهيَّأ لهذا الدور بعينه", "وما يُهيَّأ لهذا الدور بعينه"), SETUP,
    bg=TEALBG, line=TEAL)

# ═════════ المشاهد ═════════
heading("المشاهد", "وما بين قوسين زمنُ المشهد في الشريط")
t0 = 0


def scene(num, secs, title, says, img=None, extra=None):
    """⛔ **كان المشهدُ يكرّر نصَّه مرّتين**: «تفعل» ثم «تقول» بالنصّ نفسِه —
    فيقرأ المسجِّلُ سطراً ويظنّ الثانيَ غيرَه. فصار كتلةً واحدة: ما يُقال
    **وأنت تنقر**، ومعه اسمُ الشاشة فيعرفها قبل أن يفتحها."""
    global t0
    p = para(6, 2, keep=True)
    par_shd(p, NUMBG)
    run(p, "  مشهد %s  " % arn(num), size=10, bold=True, color=RED)
    run(p, "  %s – %s  " % (mmss(t0), mmss(t0 + secs)), size=8.6, color=TEAL, bold=True)
    run(p, "   " + title, size=10.4, bold=True, color=NAVY)
    t0 += secs
    if img:
        p = para(1, 1)
        _add_pPr_el(p, 'ind', {'right': '220'})
        run(p, "الشاشة:  ", size=8.8, bold=True, color=TEAL)
        run(p, V.SCREEN.get(img, img), size=9, color=NAVY)
        shot(img)
    if extra:
        for d in extra:
            bullet(d, size=9.2)
    p = para(2, 1)
    _add_pPr_el(p, 'ind', {'right': '220'})
    par_shd(p, SOFT)
    run(p, G("تقول وأنت تنقر:", "تقولين وأنتِ تنقرين:"), size=9, bold=True, color=RED)
    for x in (says if isinstance(says, (list, tuple)) else [says]):
        pp = para(0, 2)
        _add_pPr_el(pp, 'ind', {'right': '400'})
        run(pp, "«  ", size=9, color=TEAL)
        runmd(pp, x, size=9.6, color=INK)
        run(pp, "  »", size=9, color=TEAL)


scene(1, V.T_OPEN, G("الافتتاح — وأنت على شاشة اختيار الدور",
                     "الافتتاح — وأنتِ على شاشة اختيار الدور"),
      OPEN_SAY, "login",
      extra=[G("افتح الرابط وأنت صامت، ثم ابدأ الكلام والشاشةُ ظاهرة:  ",
               "افتحي الرابط وأنتِ صامتة، ثم ابدئي الكلام والشاشةُ ظاهرة:  ") + C.SITE])

for i, (title, img, lines, _w) in enumerate(S["steps"], 2):
    scene(i, V.T_STEP, title, lines, img)

scene(len(S["steps"]) + 2, V.T_CLOSE, G("الختام", "الختام"), CLOSE_SAY)

# ═════════ مصائدُ التسجيل ═════════
heading(G("ما يُفسد التسجيل", "ما يُفسد التسجيل"),
        "أعِد المشهدَ ولا تُكمل عليه — فالمشاهدُ يرى ما لا تراه")
box(G("مصائدُ هذا الدور — قُلها في الفيديو ولا تقع فيها",
      "مصائدُ هذا الدور — قوليها في الفيديو ولا تقعي فيها"),
    ["**%s** — %s" % (a, b) for a, b in S["traps"]], bg=WARNBG, line=WARNLN)
box(G("ومصائدُ التسجيل نفسِه", "ومصائدُ التسجيل نفسِه"), [
    "**اسمٌ حقيقيٌّ على الشاشة** — أوقف التسجيل وامسح البيانات وابدأ من جديد.",
    "**نقرةٌ بلا كلام** — المشاهدُ يرى الشاشةَ تتغيّر ولا يعرف لماذا.",
    "**الشرحُ قبل النقر بثوانٍ** — قُل وأنت تنقر، لا قبلها ولا بعدها.",
    "**فتحُ شاشةٍ ليست لهذا الدور** — تُربك المشاهدَ وتُوهمه بصلاحيةٍ ليست له.",
    "**رابطُ المخزن المشترك ظاهراً** — فيه مفتاحُ القاعدة، فلا يُصوَّر.",
], bg=WARNBG, line=WARNLN)

# ═════════ الخاتمة ═════════
p = para(6, 0)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, C.SITE, size=9.6, bold=True, color=TEAL, latin=True)
p = para(0, 0)
_add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "إدارة التخطيط والاعتماد المدرسي  ·  أ. أحمد صيام", size=8.4, color=GREY, light=True)

normalize_tables(doc)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("حُفظ: %s | %d مشهداً | %s دقيقة"
      % (os.path.basename(OUT), len(S["steps"]) + 2, mmss(TOTAL)))
