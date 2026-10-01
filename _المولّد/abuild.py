# -*- coding: utf-8 -*-
"""بطاقة الاتجاه التدريسي — صفحة واحدة لكل اتجاه، بهوية استمارة الملاحظة الصفية.
الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 abuild.py <مفتاح الاتجاه> <الملف>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from dox import *
from dox import GENDER
from dox import _add_pPr_el, _trPr_set, _tcPr_set, _tblPr_set, PPR_ORDER, TCPR_ORDER, _insert_ordered
import acontent as A

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, HEADBG, NUMBG, ZEBRA, INNER = "C00000", "D9D9D9", "E9EEF6", "F6F8FB", "C9D2DE"
GREY, INK, ROSE = "7F7F7F", "1A1A1A", "FCEEEE"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
ARQ = "٠١٢٣٤٥٦٧٨٩"
def ar(n): return "".join(ARQ[int(c)] if c.isdigit() else c for c in str(n))

doc = Document(); set_doc_defaults(doc); set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.0)
s.top_margin, s.bottom_margin = Cm(2.85), Cm(2.7)
s.header_distance, s.footer_distance = Cm(0.2), Cm(1.6)
s._sectPr.append(OxmlElement('w:bidi'))

hp = s.header.paragraphs[0]; hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run(); r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=16, inner=INNER, isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def table(rows, widths, **kw):
    return frame(make_table(doc, rows, widths, header_rows=0), **kw)


def span(row, a, b, widths):
    c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
    _tcPr_set(c, 'tcW', {'w': str(int(sum(widths[a:b + 1]) / sum(widths) * 5000)), 'type': 'pct'})
    return c


def auto_line(p, mult=1.0):
    _add_pPr_el(p, 'spacing', {'after': '0', 'line': str(int(240 * mult)), 'lineRule': 'auto'})
    return p


def heading(text, hint=None, sb=5):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, sb, 2); par_keep(p)
    par_border_bottom(p, TEAL, "8")
    run(p, "❖  " + text, size=10.5, bold=True, color=RED)
    if hint: run(p, "    " + hint, size=7.8, color=GREY, light=True)
    return p


def bullets(items, size=9.2, num=True, color=INK):
    """قائمة مرقّمة أو منقّطة في المتن."""
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph(); rtl_par(p); auto_line(p); par_space(p, 0, 4)
        _add_pPr_el(p, 'ind', {'right': '170'})
        run(p, (ar(i) + "  ") if num else "•  ", size=size, bold=True, color=TEAL)
        run(p, t, size=size, color=color)


# ═════════ الاتجاه المطلوب ═════════
KEY = sys.argv[1] if len(sys.argv) > 1 else "afl"
key, NAME, DEF, WHY, STEPS, EVID, EXAMPLES, NOT_, INDS, STRATS, SOURCE = \
    next(a for a in A.APPROACHES if a[0] == KEY)

if GENDER == "f":     # الصيغة المؤنثة اليدوية لما فيه فعل أو ضمير يعود على مؤنث
    fx = lambda t: A.FEM.get(t, t)
    DEF, WHY, SOURCE = fx(DEF), fx(WHY), fx(SOURCE)
    STEPS = [fx(t) for t in STEPS]; EVID = [fx(t) for t in EVID]; NOT_ = [fx(t) for t in NOT_]
    EXAMPLES = [(sub, fx(ex)) for sub, ex in EXAMPLES]

from band import band_title
band_title(s, "الاتجاه التدريسي")
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 3); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED); run(p, NAME, size=11, bold=True, color=RED); run(p, "  ◆", size=7, color=RED)

# التعريف ولماذا
t = table(2, [2.6, 16.4])
for i, (lab, txt) in enumerate((("ما هو؟", DEF), ("لماذا؟", WHY))):
    r = t.rows[i]; row_nosplit(r)
    cell_text(r.cells[0], lab, size=9, bold=True, color=NAVY, shd=NUMBG, align='center', valign='center')
    auto_line(cell_text(r.cells[1], txt, size=9.4, color=INK, valign='center', shd=ZEBRA if i else None))

heading("كيف يظهر في الحصة؟", "أربع خطوات عملية", sb=4)
bullets(STEPS)

heading("ثلاثة شواهد يراها الزائر" if GENDER == "m" else "ثلاثة شواهد تراها الزائرة", "إن غابت فالاتجاه لم يُطبَّق")
t = table(1, [W / 3] * 3)
row_nosplit(t.rows[0])
for i, e in enumerate(EVID):
    auto_line(cell_text(t.rows[0].cells[i], "✓  " + e, size=8.8, color=INK, valign='top', shd=TEALBG))

heading("أمثلة من كل مادة", "المثال مقتضب ليُحتذى لا ليُنقل")
t = table(len(EXAMPLES), [4.8, 14.2])
for i, (sub, ex) in enumerate(EXAMPLES):
    r = t.rows[i]; row_nosplit(r)
    cell_text(r.cells[0], sub, size=8.8, bold=True, color=NAVY, shd=NUMBG, valign='center')
    auto_line(cell_text(r.cells[1], ex, size=9.0, color=INK, valign='center', shd=ZEBRA if i % 2 else None))

heading("لا يُعدّ تطبيقاً للاتجاه")
t = table(1, [W / 3] * 3)
row_nosplit(t.rows[0])
for i, e in enumerate(NOT_):
    auto_line(cell_text(t.rows[0].cells[i], "✗  " + e, size=8.8, color=RED, valign='top', shd=ROSE))

# الربط بالاستمارة والبنك
t = table(2, [3.4, 15.6])
for i, (lab, txt) in enumerate((("مؤشرات تقوّيها", INDS), ("استراتيجيات تخدمه", STRATS))):
    r = t.rows[i]; row_nosplit(r)
    cell_text(r.cells[0], lab, size=8.8, bold=True, color=NAVY, shd=NUMBG, align='center', valign='center')
    auto_line(cell_text(r.cells[1], txt, size=8.8, bold=True, color=TEAL, valign='center', shd=TEALBG))

p = doc.add_paragraph(); rtl_par(p); par_space(p, 3, 0, 10); _add_pPr_el(p, 'jc', {'val': 'both'})
run(p, "السند: ", size=8, bold=True, color=GREY); run(p, SOURCE, size=8, color=GREY, light=True)

_p = doc.add_paragraph(); rtl_par(_p); par_space(_p, 0, 0, 1); run(_p, "", size=1)

fp = s.footer.paragraphs[0]; fp.text = ""
rtl_par(fp); _add_pPr_el(fp, 'jc', {'val': 'center'}); par_space(fp, 0, 0)
run(fp, "الاتجاه التدريسي — " + NAME + "   ·   مدارس ابن خلدون", size=7, color=GREY, light=True)

normalize_tables(doc)
out = sys.argv[2] if len(sys.argv) > 2 else "a.docx"
doc.save(out); print("أمثلة:", len(EXAMPLES), "| حُفظ:", out)
