# -*- coding: utf-8 -*-
"""بطاقة تشخيص أداء المعلّم في إستراتيجية تدريس — بهوية استمارة الملاحظة الصفية.
الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 sbuild.py <مفتاح الاستراتيجية> <الملف>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from dox import *
from dox import GENDER
from dox import _add_pPr_el, _trPr_set, _tcPr_set, _tblPr_set, PPR_ORDER, _insert_ordered
import scontent as S

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, HEADBG, NUMBG, ZEBRA, INNER = "C00000", "D9D9D9", "E9EEF6", "F6F8FB", "C9D2DE"
GREY, INK = "7F7F7F", "1A1A1A"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
ARQ = "٠١٢٣٤٥٦٧٨٩"
def ar(n): return "".join(ARQ[int(c)] if c.isdigit() else c for c in str(n))

doc = Document(); set_doc_defaults(doc); set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.0)
s.top_margin, s.bottom_margin = Cm(2.85), Cm(2.7)
s.header_distance, s.footer_distance = Cm(0.2), Cm(1.6)
s._sectPr.append(OxmlElement('w:bidi'))

# ═════════ الكليشة خلفيةً ═════════
hp = s.header.paragraphs[0]; hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run(); r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=16, inner=INNER, isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def table(rows, widths, **kw):
    return frame(make_table(doc, rows, widths, header_rows=0), **kw)


def span(row, a, b, widths):
    c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
    _tcPr_set(c, 'tcW', {'w': str(int(sum(widths[a:b + 1]) / sum(widths) * 5000)), 'type': 'pct'})
    return c


def dotted_lines(cell, width_cm, n=3, size=7.8):
    """أسطر كتابة بنقاط تمتد إلى نهاية الخانة (علامة جدولة بقائد نقطي)."""
    end = int(width_cm * 566.93) - 140 - 30
    cell.text = ""
    for i in range(n):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        rtl_par(p)
        tabs = OxmlElement('w:tabs')
        tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'left'); tb.set(qn('w:leader'), 'dot')
        tb.set(qn('w:pos'), str(end)); tabs.append(tb)
        _insert_ordered(p._p.get_or_add_pPr(), tabs, PPR_ORDER)
        _add_pPr_el(p, 'spacing', {'after': '0', 'line': '360', 'lineRule': 'auto'})   # ١٫٥ سطر
        run(p, ARQ[i + 1] + " ", size=size, color=GREY)
        run(p, "", size=size, color=GREY)._r.append(OxmlElement('w:tab'))
    cell_valign(cell, 'center')



def auto_line(p, mult=1.0):
    """سطر تلقائي مفرد (أو ١٫٥) بدل الارتفاع المضبوط — تعديل المستشار 2026-09-27."""
    _add_pPr_el(p, 'spacing', {'after': '0', 'line': str(int(240 * mult)), 'lineRule': 'auto'})


def cell_border(cell, sides, sz=8, color=NAVY):
    """حدّ داكن على جهة من الخلية (يُبقي بقية الحدود على حدّ الجدول الرفيع)."""
    tcPr = cell._tc.get_or_add_tcPr()
    bd = tcPr.find(qn('w:tcBorders'))
    if bd is None:
        bd = OxmlElement('w:tcBorders')
        from dox import TCPR_ORDER
        _insert_ordered(tcPr, bd, TCPR_ORDER)
    for side in sides:
        for old in bd.findall(qn('w:' + side)): bd.remove(old)
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        bd.append(e)



def sep(row):
    """فاصلان كحليان: بين «مؤشر الأداء» والمستويات، وبين المستويات و«ملاحظات»."""
    n = len(row.cells)
    cell_border(row.cells[1], ['right'])
    cell_border(row.cells[2], ['left'])
    cell_border(row.cells[n - 2], ['right'])
    cell_border(row.cells[n - 1], ['left'])


# ═════════ الاستراتيجية المطلوبة ═════════
KEY = sys.argv[1] if len(sys.argv) > 1 else "jigsaw"
NAME, DESC, INDS = next((n, d, i) for k, n, d, i in S.BANK if k == KEY)
if GENDER == "f":                      # الصيغة المؤنثة اليدوية لما فيه ضمير عائد
    INDS = [S.FEM.get(t, t) for t in INDS]

from band import band_title
band_title(s, "بطاقة تشخيص أداء المعلّم")
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 4); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED)
run(p, "إستراتيجية التعلم: " + NAME, size=10.5, bold=True, color=RED)
run(p, "  ◆", size=7, color=RED)

# ═════════ بيانات الحصة ═════════
DW = [1.9, 2.4, 2.45, 1.6, 1.5, 1.5, 2.0, 1.5, 2.1, 2.05]
t = table(2, DW)
for ri, labs in enumerate([("اسم المعلم", "المرحلة", "المادة", "الصف/الفصل", "الحصة"),
                           ("موضوع الدرس", "الاتجاه التدريسي", "التاريخ", "اليوم", "زمن الحصة")]):
    row_height(t.rows[ri], 0.52); row_nosplit(t.rows[ri])
    for j, lb in enumerate(labs):
        cell_text(t.rows[ri].cells[j * 2], lb, size=8, bold=True, color=NAVY, shd=NUMBG, align='center')
        cell_text(t.rows[ri].cells[j * 2 + 1], "", size=8)
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 4)

# ═════════ مصفوفة المؤشرات ═════════
MW = [0.75, 6.95, 1.5, 1.5, 1.5, 1.5, 1.95, 3.35]   # تعديل المستشار 2026-09-27: ملاحظات أوسع
t = table(2, MW)
h = t.rows[0]; row_height(h, 0.44); row_nosplit(h); _trPr_set(h, 'tblHeader')
cell_text(h.cells[0], "م", size=8.5, bold=True, color=NAVY, shd=HEADBG, align='center')
cell_text(h.cells[1], "مؤشر الأداء", size=8.5, bold=True, color=NAVY, shd=HEADBG, align='center')
c = span(h, 2, 6, MW); cell_text(c, "مستوى الأداء", size=8.2, bold=True, color=NAVY, shd=HEADBG, align='center')
cell_text(h.cells[7], "ملاحظات", size=8.2, bold=True, color=NAVY, shd=HEADBG, align='center')
sep(h)
cell_border(h.cells[2], ['top'])
h2 = t.rows[1]; row_height(h2, 0.5); row_nosplit(h2); _trPr_set(h2, 'tblHeader')
cell_text(h2.cells[0], "", size=5, shd=HEADBG); cell_text(h2.cells[1], "", size=5, shd=HEADBG)
for k, (lab, deg) in enumerate(S.LEVELS, 2):
    cl = h2.cells[k]; cl.text = ""
    pp = cl.paragraphs[0]; rtl_par(pp); par_space(pp, 0, 0, 9); _add_pPr_el(pp, 'jc', {'val': 'center'})
    run(pp, lab + " " + deg, size=6.5, bold=True, color=NAVY)   # تعديل المستشار: مقاس واحد أصغر
    cell_shd(cl, NUMBG); cell_valign(cl, 'center')
cell_text(h2.cells[7], "", size=5, shd=HEADBG)
sep(h2)


def fixw(row):
    for i, w in enumerate(MW):
        _tcPr_set(row.cells[i], 'tcW', {'w': str(int(w / sum(MW) * 5000)), 'type': 'pct'})


for n, txt in enumerate(INDS, 1):
    r = t.add_row(); fixw(r); row_nosplit(r)   # بلا ارتفاع ثابت: الصف يتبع نصّه (تعديل المستشار)
    z = ZEBRA if n % 2 == 0 else None
    auto_line(cell_text(r.cells[0], ar(n), size=8, bold=True, color=NAVY, align='center', shd=NUMBG))
    auto_line(cell_text(r.cells[1], txt, size=8.3, color=INK, valign='center', shd=z))
    for k in range(2, 8):
        auto_line(cell_text(r.cells[k], "", size=5, align='center', shd=z))
    sep(r)
    if n == len(INDS):     # صف المؤشر الأخير: خط كحلي يفصله عن صف المجموع، وإطار جانبي
        for k in range(8): cell_border(r.cells[k], ['bottom'])
        cell_border(r.cells[0], ['left']); cell_border(r.cells[7], ['right'])

r = t.add_row(); fixw(r); row_height(r, 0.5); row_nosplit(r)
c = span(r, 0, 1, MW); c.text = ""
pp = c.paragraphs[0]; rtl_par(pp); par_space(pp, 0, 0, 12); _add_pPr_el(pp, 'jc', {'val': 'center'})
run(pp, "مجموع البطاقة", size=9, bold=True, color=TEAL)
run(pp, "   (عشرة مؤشرات × ١٠)", size=7.3, color=GREY, light=True)
cell_shd(c, TEALBG); cell_valign(c, 'center')
c = span(r, 2, 5, MW)
cell_text(c, "........  من ١٠٠", size=8.6, bold=True, color=TEAL, align='center', shd=TEALBG)
c2 = span(r, 6, 7, MW)
cell_text(c2, "النسبة: ........ ٪", size=8.2, bold=True, color=TEAL, align='center', shd=TEALBG)
for _c in (r.cells[0], c, c2): cell_border(_c, ['top'])

# ═════════ اشتقاق درجة الاستمارة ═════════
p = doc.add_paragraph(); rtl_par(p); par_space(p, 3, 3, 11); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "تُشتقّ منها درجة المؤشر م٢ · ٣ في استمارة الملاحظة الصفية:  ", size=8.3, bold=True, color=NAVY)
run(p, S.CONVERT, size=8, color=INK)

# ═════════ إيجابيات · نقاط للتأمل ═════════
NW2 = [4.4, 14.6]
t = table(2, NW2)
for ri, lab in enumerate(["إيجابيات", "نقاط للتأمل"]):
    r = t.rows[ri]; row_height(r, 2.2); row_nosplit(r)
    cell_text(r.cells[0], lab, size=9, bold=True, color=NAVY, shd=NUMBG, valign='center', align='center')
    dotted_lines(r.cells[1], NW2[1], n=5)
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 4)

# ═════════ البرنامج العلاجي ═════════
RW = [3.4] + [(W - 3.4) / 3] * 3
t = table(2, RW)
c = t.rows[0].cells[0].merge(t.rows[1].cells[0])
_tcPr_set(c, 'tcW', {'w': str(int(3.4 / W * 5000)), 'type': 'pct'})
cell_text(c, "البرنامج العلاجي", size=9, bold=True, color=NAVY, shd=NUMBG, valign='center', align='center')
for i, item in enumerate(S.REMEDIAL):
    rr, k = divmod(i, 3)
    row_height(t.rows[rr], 0.62); row_nosplit(t.rows[rr])
    cell_text(t.rows[rr].cells[k + 1], "☐  " + item, size=8.2, color=INK, valign='center',
              shd=ZEBRA if rr else None)
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 4)

# ═════════ التوقيعات ═════════
GW = [2.3, 3.2, 1.5, 2.2, 4.2, 2.2, 1.5, 1.9]
t = table(1, GW)
row_height(t.rows[0], 1.04); row_nosplit(t.rows[0])
for i3, lb in enumerate(["المعلم المنفذ", "", "التوقيع", "", "الزائر: المدير / الوكيل / المشرف", "", "التوقيع", ""]):
    cell_text(t.rows[0].cells[i3], lb, size=7.6, bold=bool(lb), color=NAVY, shd=NUMBG if lb else None,
              align='center')

_p = doc.add_paragraph(); rtl_par(_p); par_space(_p, 0, 0, 1); run(_p, "", size=1)

# ═════════ التذييل ═════════
fp = s.footer.paragraphs[0]; fp.text = ""
rtl_par(fp); _add_pPr_el(fp, 'jc', {'val': 'center'}); par_space(fp, 0, 0)
run(fp, "بطاقة تشخيص أداء المعلّم — " + NAME + "   ·   مدارس ابن خلدون", size=7, color=GREY, light=True)

normalize_tables(doc)
out = sys.argv[2] if len(sys.argv) > 2 else "s.docx"
doc.save(out); print("مؤشرات:", len(INDS), "| حُفظ:", out)
