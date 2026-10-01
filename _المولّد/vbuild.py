# -*- coding: utf-8 -*-
"""بطاقة زيارة الأقران — «شاهدتُ وأطبّق»: زيارتان في الفصل الدراسي، بهوية أدوات ابن خلدون.
بديلُ ملف النمو المهني في البرنامج السابق: الزيارةُ تنتهي بإجراءٍ يُنقل إلى بطاقة الجسر.
الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 vbuild.py <الملف>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from dox import *
from dox import GENDER, _add_pPr_el, _trPr_set, _tcPr_set, _tblPr_set, PPR_ORDER, _insert_ordered

NAVY, TEAL, TEALBG = "355E91", "2F7F95", "E2F0F3"
RED, HEADBG, NUMBG, ZEBRA, INNER = "C00000", "D9D9D9", "E9EEF6", "F6F8FB", "C9D2DE"
GREY, INK = "7F7F7F", "1A1A1A"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
ARQ = "٠١٢٣٤٥٦٧٨٩"

doc = Document(); set_doc_defaults(doc); set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.0)
s.top_margin, s.bottom_margin = Cm(2.85), Cm(2.7)
s.header_distance, s.footer_distance = Cm(0.2), Cm(1.6)
s._sectPr.append(OxmlElement('w:bidi'))

hp = s.header.paragraphs[0]; hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run(); r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=16, inner=INNER, isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def table(rows, widths, **kw):
    return frame(make_table(doc, rows, widths, header_rows=0), **kw)


def span(row, a, b, widths):
    c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
    _tcPr_set(c, 'tcW', {'w': str(int(sum(widths[a:b + 1]) / sum(widths) * 5000)), 'type': 'pct'})
    return c


def auto_line(p, mult=1.0):
    _add_pPr_el(p, 'spacing', {'after': '0', 'line': str(int(240 * mult)), 'lineRule': 'auto'})
    return p


def dotted(cell, width_cm, n=3, size=8, mult=1.0):
    end = int(width_cm * 566.93) - 140 - 30
    cell.text = ""
    for i in range(n):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        rtl_par(p)
        tabs = OxmlElement('w:tabs')
        tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'left'); tb.set(qn('w:leader'), 'dot')
        tb.set(qn('w:pos'), str(end)); tabs.append(tb)
        _insert_ordered(p._p.get_or_add_pPr(), tabs, PPR_ORDER)
        _add_pPr_el(p, 'spacing', {'after': '0', 'line': str(int(240 * mult)), 'lineRule': 'auto'})
        run(p, ARQ[i + 1] + " ", size=size, color=GREY)
        run(p, "", size=size, color=GREY)._r.append(OxmlElement('w:tab'))
    cell_valign(cell, 'center')


def heading(text, hint=None, sb=5):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, sb, 2); par_keep(p)
    par_border_bottom(p, TEAL, "8")
    run(p, "❖  " + text, size=10.5, bold=True, color=RED)
    if hint: run(p, "    " + hint, size=7.8, color=GREY, light=True)
    return p


def ticks_row(cell, items, size=7.8):
    cell.text = ""
    p = cell.paragraphs[0]; rtl_par(p); auto_line(p); _add_pPr_el(p, 'jc', {'val': 'both'})
    for i, it in enumerate(items):
        run(p, "☐ " + it, size=size, color=INK)
        if i < len(items) - 1: run(p, "    ", size=size, color=INK)
    cell_valign(cell, 'center')


MAJALAT = ["التخطيط وتجهيز بيئة التعلم", "العرض وإدارة الحصة", "مشاركة المتعلمين واندماجهم",
           "المهارات المستهدفة", "التقنية والتعلم الإلكتروني", "التقويم والتغذية الراجعة",
           "بناء شخصية الطالب", "إستراتيجية التدريس", "الضبط الصفي", "الحصص المعملية"]

from band import band_title
band_title(s, "بطاقة زيارة الأقران")
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 4); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED)
run(p, "شاهدتُ وأطبّق  ·  زيارتان لزميلين في الفصل الدراسي", size=10, bold=True, color=RED)
run(p, "  ◆", size=7, color=RED)

# بيانات المعلم الزائر
DW = [2.6, 3.6, 1.6, 2.4, 1.9, 2.3, 2.6, 2.0]
t = table(1, DW)
row_height(t.rows[0], 0.55); row_nosplit(t.rows[0])
for j, lb in enumerate(["المعلم الزائر", "المادة", "المرحلة", "الفصل الدراسي"]):
    cell_text(t.rows[0].cells[j * 2], lb, size=8.2, bold=True, color=NAVY, shd=NUMBG, align='center', valign='center')
    cell_text(t.rows[0].cells[j * 2 + 1], "", size=8)

for v in (1, 2):
    heading(f"الزيارة {'الأولى' if v == 1 else 'الثانية'}",
            "تُملأ في الحصة نفسها، وتُسلَّم للوكيل التعليمي خلال يومين" if v == 1 else None,
            sb=4 if v == 1 else 3)
    # بيانات الحصة المستضافة
    VW = [2.7, 2.8, 2.5, 1.6, 1.2, 1.4, 1.2, 1.5, 1.9, 2.2]
    t = table(2, VW)
    for ri, labs in enumerate([("المعلم المستضيف", "عنوان الدرس", "الصف", "الحصة", "التاريخ"),
                               ("الإستراتيجية", "الاتجاه التدريسي", "المادة", "اليوم", "زمن الحضور")]):
        row_height(t.rows[ri], 0.48); row_nosplit(t.rows[ri])
        for j, lb in enumerate(labs):
            cell_text(t.rows[ri].cells[j * 2], lb, size=8, bold=True, color=NAVY, shd=NUMBG, align='center', valign='center')
            cell_text(t.rows[ri].cells[j * 2 + 1], "", size=8)
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 3)

    # المجال المستهدف
    t = table(2, [3.0, 16.0])
    c = t.rows[0].cells[0].merge(t.rows[1].cells[0])
    _tcPr_set(c, 'tcW', {'w': str(int(3.0 / W * 5000)), 'type': 'pct'})
    cell_text(c, "المجال المستهدف", size=8.4, bold=True, color=NAVY, shd=NUMBG, align='center', valign='center')
    for ri, part in enumerate((MAJALAT[:5], MAJALAT[5:])):
        row_height(t.rows[ri], 0.4); row_nosplit(t.rows[ri])
        ticks_row(t.rows[ri].cells[1], part)
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 3)

    # ماذا استفدت · ماذا سأطبّق — جدول صريح: رقم السطر بجوار خانة كتابته (لا علامات جدولة)
    NW = [3.0, 0.9, 10.7, 2.4, 2.0]
    t = table(5, NW)
    lab = t.rows[0].cells[0]
    for rr in range(1, 3):
        lab = lab.merge(t.rows[rr].cells[0])
    _tcPr_set(lab, 'tcW', {'w': str(int(3.0 / W * 5000)), 'type': 'pct'})
    cell_lines(lab, [("ثلاث نقاط", True), ("استفدتها في حصتي", True)], size=8.4, color=NAVY,
               shd=NUMBG, valign='center', align='center', space=1)
    for i in range(3):
        r = t.rows[i]; row_height(r, 0.62); row_nosplit(r)
        cell_text(r.cells[1], ARQ[i + 1], size=8.4, bold=True, color=NAVY, shd=NUMBG,
                  align='center', valign='center')
        c = span(r, 2, 4, NW); cell_text(c, "", size=8, valign='center')
    # صف الإجراء
    r = t.rows[3]; row_height(r, 0.66); row_nosplit(r)
    cell_lines(r.cells[0], [("إجراء واحد", True), ("سأطبّقه", True)], size=8.4, color=NAVY,
               shd=TEALBG, valign='center', align='center', space=1)
    c = span(r, 1, 2, NW); cell_text(c, "", size=8, valign='center')
    cell_text(r.cells[3], "تاريخ التطبيق", size=8, bold=True, color=NAVY, shd=TEALBG,
              align='center', valign='center')
    cell_text(r.cells[4], "", size=8, valign='center')
    # صف نقل الإجراء إلى بطاقة الجسر
    r = t.rows[4]; row_height(r, 0.46); row_nosplit(r)
    cell_text(r.cells[0], "المتابعة", size=8, bold=True, color=NAVY, shd=TEALBG,
              align='center', valign='center')
    c = span(r, 1, 4, NW)
    cell_text(c, "☐  نُقل الإجراء إلى بطاقة الجسر الخاصة بي، وسيُتحقق منه في الزيارة القادمة",
              size=8, bold=True, color=TEAL, valign='center')
    p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 3)

    # التوقيعات
    GW = [2.8, 3.6, 1.6, 2.4, 3.0, 3.2, 1.2, 1.2]
    t = table(1, GW)
    row_height(t.rows[0], 0.5); row_nosplit(t.rows[0])
    for j, lb in enumerate(["المعلم المستضيف", "", "التوقيع", "", "المعلم الزائر", "", "التوقيع", ""]):
        cell_text(t.rows[0].cells[j], lb, size=7.8, bold=bool(lb), color=NAVY,
                  shd=NUMBG if lb else None, align='center', valign='center')

heading("اعتماد المتابعة", "يراجعها الوكيل التعليمي، ويتحقق من تنفيذ الإجراء في الزيارة التالية", sb=3)
t = table(1, [3.2, 5.0, 2.0, 3.0, 2.8, 3.0])
row_height(t.rows[0], 0.58); row_nosplit(t.rows[0])
for j, lb in enumerate(["الوكيل التعليمي", "", "التوقيع", "", "تاريخ التحقق من الإجراء", ""]):
    cell_text(t.rows[0].cells[j], lb, size=7.8, bold=bool(lb), color=NAVY,
              shd=NUMBG if lb else None, align='center', valign='center')

_p = doc.add_paragraph(); rtl_par(_p); par_space(_p, 0, 0, 1); run(_p, "", size=1)

fp = s.footer.paragraphs[0]; fp.text = ""
rtl_par(fp); _add_pPr_el(fp, 'jc', {'val': 'center'}); par_space(fp, 0, 0)
run(fp, "بطاقة زيارة الأقران — مدارس ابن خلدون", size=7, color=GREY, light=True)

normalize_tables(doc)
out = sys.argv[1] if len(sys.argv) > 1 else "v.docx"
doc.save(out); print("حُفظ:", out)
