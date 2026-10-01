# -*- coding: utf-8 -*-
"""استمارة الملاحظة الصفية — كليشة ابن خلدون خلفيةً · العنوان في الشريط · إطار ٢ نقطة ·
صف مجموع لكل مجال · ثلاث صفحات. أُعيد بناء المولّد 2026-09-27 بعد ضياعه، مطابقاً للمنشور.
الاستعمال: CLS_FONTSET=js CLS_GENDER=m|f python3 zbuild2.py <الملف>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from dox import *
from dox import GENDER, _add_pPr_el, _trPr_set, _tcPr_set, _tblPr_set, PPR_ORDER, _insert_ordered
import zcontent as Z

NAVY   = "355E91"
TEAL   = "2F7F95"
TEALBG = "E2F0F3"
RED    = "C00000"
HEADBG = "D9D9D9"
NUMBG  = "E9EEF6"
ZEBRA  = "F6F8FB"
NABG   = "F2F2F2"
INNER  = "C9D2DE"
GREY   = "7F7F7F"
INK    = "1A1A1A"
W = 19.0
HERE = os.path.dirname(os.path.abspath(__file__))
KLISHA = os.path.join(HERE, "kl_portrait.jpg")
ARQ = "٠١٢٣٤٥٦٧٨٩"
def ar(n): return "".join(ARQ[int(c)] if c.isdigit() else c for c in str(n))

doc = Document(); set_doc_defaults(doc); set_compat15(doc)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.0)
s.top_margin, s.bottom_margin = Cm(2.85), Cm(2.7)
s.header_distance, s.footer_distance = Cm(0.2), Cm(1.6)
s._sectPr.append(OxmlElement('w:bidi'))

# ═════════ الكليشة خلفيةً ═════════
hp = s.header.paragraphs[0]; hp.text = ""
par_space(hp, 0, 0)
r = hp.add_run(); r.add_picture(KLISHA, width=Cm(21.0), height=Cm(29.7))
inline = r._r.find('.//' + qn('wp:inline'))
graphic = inline.find(qn('a:graphic'))
anchor = parse_xml(
    f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
    'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="7560000" cy="10692000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/><wp:docPr id="901" name="klisha"/><wp:cNvGraphicFramePr/></wp:anchor>')
anchor.append(graphic)
inline.getparent().replace(inline, anchor)


def frame(t, outer=NAVY, osz=16, inner=INNER, isz=4):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    bd = OxmlElement('w:tblBorders')
    for side, col, sz in (('top', outer, osz), ('left', outer, osz), ('bottom', outer, osz),
                          ('right', outer, osz), ('insideH', inner, isz), ('insideV', inner, isz)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        bd.append(e)
    from dox import TBLPR_ORDER
    _insert_ordered(tblPr, bd, TBLPR_ORDER)
    return t


def table(rows, widths, **kw):
    return frame(make_table(doc, rows, widths, header_rows=0), **kw)


def span(row, a, b, widths):
    c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
    _tcPr_set(c, 'tcW', {'w': str(int(sum(widths[a:b + 1]) / sum(widths) * 5000)), 'type': 'pct'})
    return c


def heading(text, hint=None, sb=5, brk=False):
    p = doc.add_paragraph(); rtl_par(p); par_space(p, sb, 2); par_keep(p)
    if brk: par_break_before(p)
    par_border_bottom(p, TEAL, "8")
    run(p, "❖  " + text, size=10.5, bold=True, color=RED)
    if hint: run(p, "    " + hint, size=7.8, color=GREY, light=True)
    return p


def dotted_lines(cell, width_cm, with_date=False, n=3, size=7.8):
    """أسطر كتابة تمتد نقاطها إلى نهاية الخانة (علامة جدولة بقائد نقطي)."""
    end = int(width_cm * 566.93) - 140 - 30
    stops = [end - 1900, end] if with_date else [end]
    cell.text = ""
    for i in range(n):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        rtl_par(p)
        tabs = OxmlElement('w:tabs')
        for pos in stops:
            tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'left'); tb.set(qn('w:leader'), 'dot')
            tb.set(qn('w:pos'), str(pos)); tabs.append(tb)
        _insert_ordered(p._p.get_or_add_pPr(), tabs, PPR_ORDER)
        par_space(p, 0, 3 if i < n - 1 else 0, round(size * LINE_K, 1))
        run(p, ARQ[i + 1] + " ", size=size, color=GREY)
        run(p, "", size=size, color=GREY)._r.append(OxmlElement('w:tab'))
        if with_date:
            run(p, "   تاريخ التحقق: ", size=size, color=GREY)
            run(p, "", size=size, color=GREY)._r.append(OxmlElement('w:tab'))
    cell_valign(cell, 'center')


# ═════════ العنوان ═════════
from band import band_title
band_title(s, "استمارة الملاحظة الصفية")
p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 3); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "◆  ", size=7, color=RED)
run(p, "المعيار الموحَّد لتقييم الحصة الدراسية", size=9.5, bold=True, color=RED)
run(p, "  ◆", size=7, color=RED)

# ═════════ بيانات الحصة ═════════
DW = [2.0, 2.6, 1.9, 2.1, 1.5, 1.8, 2.0, 1.9, 1.9, 1.3]
t = table(2, DW)
for ri, labs in enumerate([("اسم المعلم", "المرحلة", "المادة", "الصف/الشعبة", "الحصة"),
                           ("موضوع الدرس", "الإستراتيجية", "التاريخ", "زمن الحصة", "عدد الطلاب")]):
    row_height(t.rows[ri], 0.46); row_nosplit(t.rows[ri])
    for j, lb in enumerate(labs):
        cell_text(t.rows[ri].cells[j * 2], lb, size=8, bold=True, color=NAVY, shd=NUMBG, align='center')
        cell_text(t.rows[ri].cells[j * 2 + 1], "", size=9)

p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 3)

VW = [2.3, 6.3, 10.4]
t = table(1, VW)
row_height(t.rows[0], 0.7); row_nosplit(t.rows[0])
cell_lines(t.rows[0].cells[0], [("نوع الزيارة", True), ("والاتجاه التدريسي", True)], size=8,
           color=NAVY, shd=NUMBG, valign='center', align='center', space=1)
cell_lines(t.rows[0].cells[1],
           [("☐ استكشافية     ☐ متابعة     ☐ تقويمية", False),
            ("الاتجاه التدريسي: ....................................", False)],
           size=8.3, color=INK, valign='center', align='center', space=2)
cell_lines(t.rows[0].cells[2],
           [("٤ بشواهد كاملة   ·   ٣ بشواهد كافية   ·   ٢ متحقق جزئياً   ·   ١ غير متحقق", True),
            ("«لا ينطبق»: يُطرح ٤ من المقام لكل مؤشر لا ينطبق على طبيعة الحصة", False)],
           size=7.8, color=RED, shd=NABG, valign='center', align='center', space=1)

p = doc.add_paragraph(); rtl_par(p); par_space(p, 0, 0, 3)

# ═════════ مصفوفة المؤشرات ═════════
IND_LINE = 10.2      # سطر المؤشر: أضيق قليلاً ليتسع المجال ٤ في الصفحة الأولى (طلب المستشار)
MW = [0.75, 13.6, 0.85, 0.85, 0.85, 0.85, 1.25]
t = table(1, MW)
h = t.rows[0]; row_height(h, 0.5); row_nosplit(h); _trPr_set(h, 'tblHeader')
cell_text(h.cells[0], "م", size=8.5, bold=True, color=NAVY, shd=HEADBG, align='center')
cell_text(h.cells[1], "مؤشر الأداء", size=8.5, bold=True, color=NAVY, shd=HEADBG, align='center')
c = span(h, 2, 5, MW); cell_text(c, "مستوى التحقق", size=8.2, bold=True, color=NAVY, shd=HEADBG, align='center')
cell_text(h.cells[6], "لا ينطبق", size=6.2, bold=True, color=NAVY, shd=HEADBG, align='center')



def row_keep(row):
    """إبقاء الصف مع الذي يليه — فلا ينقسم المجال بين صفحتين (طلب المستشار 2026-09-27)."""
    for c in row.cells:
        for p in c.paragraphs:
            par_keep(p)


def fixw(row):
    for i, w in enumerate(MW):
        _tcPr_set(row.cells[i], 'tcW', {'w': str(int(w / sum(MW) * 5000)), 'type': 'pct'})


for mi, (name, items) in enumerate(Z.MAJALAT, 1):
    total = len(items) * 4
    base, _, note = name.partition("  —  ")
    r = t.add_row(); fixw(r); row_height(r, 0.46); row_nosplit(r)
    c = span(r, 0, 1, MW); c.text = ""
    pp = c.paragraphs[0]; rtl_par(pp); par_space(pp, 0, 0, 12)
    run(pp, f"المجال {ar(mi)}  ·  {base}", size=9.3, bold=True, color="FFFFFF")
    if note: run(pp, f"   — {note}", size=7.5, color="DCE6F1", light=True)
    cell_shd(c, NAVY); cell_valign(c, 'center')
    for k, lb in zip(range(2, 7), ["٤", "٣", "٢", "١", "لا ينطبق"]):
        cell_text(r.cells[k], lb, size=6.2 if k == 6 else 8, bold=True, color="FFFFFF", align='center', shd=NAVY)
    row_keep(r)

    for n, (txt, _tags) in enumerate(items, 1):
        r = t.add_row(); fixw(r); row_height(r, 0.38); row_nosplit(r)
        z = ZEBRA if n % 2 == 0 else None
        cell_text(r.cells[0], ar(n), size=8, bold=True, color=NAVY, align='center', shd=NUMBG)
        cell_text(r.cells[1], txt, size=8.3, color=INK, valign='center', shd=z, line=IND_LINE)
        for k in range(2, 6):
            cell_text(r.cells[k], "", size=5, align='center', shd=z)
        cell_text(r.cells[6], "", size=5, align='center', shd=NABG)
        row_keep(r)

    r = t.add_row(); fixw(r); row_height(r, 0.44); row_nosplit(r)
    c = span(r, 0, 1, MW); c.text = ""
    pp = c.paragraphs[0]; rtl_par(pp); par_space(pp, 0, 0, 12)
    run(pp, f"مجموع المجال {ar(mi)}", size=8.8, bold=True, color=TEAL)
    cnt = len(items)
    word = "مؤشر" if cnt == 1 else ("مؤشران" if cnt == 2 else f"{ar(cnt)} مؤشرات")
    run(pp, f"    ({word} × ٤)", size=7.3, color=GREY, light=True)
    cell_shd(c, TEALBG); cell_valign(c, 'center')
    c = span(r, 2, 5, MW)
    cell_text(c, f"........  من {ar(total)}", size=8.3, bold=True, color=TEAL, align='center', shd=TEALBG)
    cell_text(r.cells[6], "", size=5, align='center', shd=TEALBG)

# ═════════ ما قاله الطلاب ═════════
heading("ما قاله الطلاب", "يُسأل ثلاثة طلاب من مستويات مختلفة — أصدق ما يكشف رسوخ الممارسة")
SW = [8.2, 3.6, 3.6, 3.6]
t = table(1 + len(Z.TULAB), SW)
for i, lb in enumerate(["السؤال", "الطالب الأول", "الطالب الثاني", "الطالب الثالث"]):
    cell_text(t.rows[0].cells[i], lb, size=8, bold=True, color=NAVY, align='center', shd=HEADBG)
row_height(t.rows[0], 0.44); row_nosplit(t.rows[0])
for i, q in enumerate(Z.TULAB, 1):
    r = t.rows[i]; row_height(r, 0.6); row_nosplit(r)
    cell_text(r.cells[0], q, size=7.8, bold=True, color=NAVY, valign='center', shd=NUMBG)
    for k in range(1, 4): cell_text(r.cells[k], "", size=8, valign='top')

# ═════════ الأثر المادي ═════════
heading("الأثر المادي في الصف", "ما يراه عضو التقويم الخارجي قبل أن ينطق المعلم")
t = table(2, [4.75] * 4)
for i, item in enumerate(Z.ATHAR):
    r, c = divmod(i, 4)
    row_height(t.rows[r], 0.44); row_nosplit(t.rows[r])
    cell_text(t.rows[r].cells[c], "☐  " + item, size=7.8, color=INK, valign='center', shd=ZEBRA if r else None)

# ═════════ الوجه الثالث ═════════
heading("بعد الحصة  ·  التحليل والتغذية الراجعة", "تُملأ في جلسة التغذية الراجعة مع المعلم", sb=4)

heading("شواهد الحكم", "ضع ✓ أمام كل مصدر بنى عليه الزائر تقديره", sb=2)
t = table(2, [4.75] * 4)
for i, item in enumerate(Z.SHAWAHED):
    r, c = divmod(i, 4)
    row_height(t.rows[r], 0.46); row_nosplit(t.rows[r])
    cell_text(t.rows[r].cells[c], "☐  " + item, size=7.6, color=NAVY, valign='center',
              shd=ZEBRA if r else None)

heading("مخطط عدالة تفاعل الطلاب", "تُرسم المقاعد، وتُوضع علامة عند كل طالب شارك — فتظهر الفجوات بالعين لا بالظن")
t = table(4, [3.8] * 5, outer=TEAL, osz=12, inner="D5E6EA")
for r in t.rows:
    row_height(r, 0.95); row_nosplit(r)
    for c in r.cells: cell_text(c, "", size=9)
p = doc.add_paragraph(); rtl_par(p); par_space(p, 1, 2, 10); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "▲  مقدمة الصف  (السبورة)", size=7.5, color=GREY, light=True)

NW = [4.4, 14.6]
t = table(2, NW)
for ri, lab in enumerate(["جوانب إبداعية\nومميزة لدى المعلم", "إجراءات تم الاتفاق عليها\nيُتحقق منها في الزيارة القادمة"]):
    r = t.rows[ri]; row_height(r, 1.75); row_nosplit(r)
    cell_lines(r.cells[0], [(x, True) for x in lab.split("\n")], size=8, color=NAVY,
               shd=NUMBG, valign='center', align='center', space=1)
    dotted_lines(r.cells[1], NW[1], with_date=(ri == 1))

heading("مجالات الدعم وترشيح الزيارة القادمة")
t = table(3, [4.75] * 4)
names8 = [m[0].split("  —  ")[0] for m in Z.MAJALAT]
for i2, nm in enumerate(names8):
    r, c = divmod(i2, 4)
    row_height(t.rows[r], 0.44); row_nosplit(t.rows[r])
    cell_text(t.rows[r].cells[c], "☐  " + nm, size=7.6, color=INK, valign='center')
row_height(t.rows[2], 0.56); row_nosplit(t.rows[2])
for c, lb in enumerate(["المعلم المُرشَّح لزيارة قادمة", "", "تاريخ الزيارة القادمة", ""]):
    cell_text(t.rows[2].cells[c], lb, size=7.6, bold=bool(lb), color=NAVY, shd=NUMBG if lb else None, align='center')

# ═════════ النتيجة ═════════
heading("النتيجة")
names = ["التخطيط", "العرض والإدارة", "المشاركة", "المهارات", "التقنية", "التقويم", "بناء الشخصية",
         "مهارات المعلم", "المجموع", "النسبة"]
tot = sum(len(m[1]) for m in Z.MAJALAT) * 4
maxes = [ar(len(m[1]) * 4) for m in Z.MAJALAT] + [ar(tot), "٪"]
t = table(3, [1.9] * 10)
for i3, nm in enumerate(names):
    cell_lines(t.rows[0].cells[i3], [(x, True) for x in nm.split("\n")], size=7.4,
               color="FFFFFF" if i3 >= 8 else NAVY, shd=NAVY if i3 >= 8 else HEADBG,
               valign='center', align='center', space=0)
    cell_text(t.rows[1].cells[i3], f"من {maxes[i3]}" if maxes[i3] != "٪" else "٪", size=7.8, bold=True,
              color=GREY, align='center', shd=NABG)
    cell_text(t.rows[2].cells[i3], "", size=10, align='center', shd=TEALBG if i3 >= 8 else None)
row_height(t.rows[0], 0.46); row_height(t.rows[1], 0.36); row_height(t.rows[2], 0.62)
for rr in t.rows: row_nosplit(rr)

p = doc.add_paragraph(); rtl_par(p); par_space(p, 2, 4, 11); _add_pPr_el(p, 'jc', {'val': 'center'})
run(p, "النسبة المئوية  =  المجموع ÷ ٢", size=8.5, bold=True, color=TEAL)
run(p, "              عدد مؤشرات «لا ينطبق»: ", size=8, bold=True, color=NAVY)
run(p, "........", size=8, color=GREY)
run(p, "   (إن وُجدت فالمقام ٢٠٠ ناقص ٤ لكل مؤشر منها)", size=7.8, color=GREY, light=True)

# سطر بطاقة الاستراتيجية: يُملأ من بطاقة التشخيص، وتُشتقّ منه درجة المؤشر م٢ · ٣
SGW = [4.3, 2.4, 2.2, 2.5, 7.6]
t = table(1, SGW)
row_height(t.rows[0], 0.56); row_nosplit(t.rows[0])
cell_text(t.rows[0].cells[0], "بطاقة تشخيص الاستراتيجية", size=8, bold=True, color=NAVY,
          shd=NUMBG, align='center', valign='center')
cell_text(t.rows[0].cells[1], "........  من ١٠٠", size=8.2, bold=True, color=TEAL,
          align='center', valign='center', shd=TEALBG)
cell_text(t.rows[0].cells[2], "النسبة: ....٪", size=8, bold=True, color=TEAL,
          align='center', valign='center', shd=TEALBG)
cell_text(t.rows[0].cells[3], "درجة م٢ · ٣: ....", size=8, bold=True, color=NAVY,
          align='center', valign='center', shd=NUMBG)
cell_text(t.rows[0].cells[4], "٩٠٪ فأعلى ← ٤   ·   ٧٥ إلى أقل من ٩٠ ← ٣   ·   ٥٠ إلى أقل من ٧٥ ← ٢   ·   أقل من ٥٠ ← ١",
          size=7.4, color=GREY, align='center', valign='center')
p = doc.add_paragraph(); rtl_par(p); par_space(p, 2, 0, 3)

GW = [1.8, 3.9, 1.5, 2.4, 4.4, 3.2, 1.8]
t = table(1, GW)
row_height(t.rows[0], 0.62); row_nosplit(t.rows[0])
for i3, lb in enumerate(["اسم المعلم", "", "التوقيع", "", "الزائر: المدير / الوكيل / المشرف", "", "التوقيع"]):
    cell_text(t.rows[0].cells[i3], lb, size=7.5, bold=bool(lb), color=NAVY, shd=NUMBG if lb else None, align='center')

_p = doc.add_paragraph(); rtl_par(_p); par_space(_p, 0, 0, 1); run(_p, "", size=1)

# ═════════ التذييل ═════════
fp = s.footer.paragraphs[0]; fp.text = ""
rtl_par(fp); _add_pPr_el(fp, 'jc', {'val': 'center'}); par_space(fp, 0, 0)
run(fp, "استمارة الملاحظة الصفية — مدارس ابن خلدون   ·   صفحة ", size=7, color=GREY, light=True)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), ' PAGE ')
fr = OxmlElement('w:r'); frpr = OxmlElement('w:rPr'); rf = OxmlElement('w:rFonts')
for a in ('ascii', 'hAnsi', 'cs'): rf.set(qn('w:' + a), FONT)
frpr.append(rf)
for tg, v in (('sz', '14'), ('szCs', '14'), ('color', GREY)):
    e = OxmlElement('w:' + tg); e.set(qn('w:val'), v); frpr.append(e)
fr.append(frpr); tt = OxmlElement('w:t'); tt.text = "1"; fr.append(tt)
fld.append(fr); fp._p.append(fld)

normalize_tables(doc)
out = sys.argv[1] if len(sys.argv) > 1 else "z2.docx"
doc.save(out)
print("مؤشرات:", sum(len(m[1]) for m in Z.MAJALAT), "| الدرجة:", tot, "| حُفظ:", out)
