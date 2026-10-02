# -*- coding: utf-8 -*-
"""⛔ حقنُ قسم **المجال ٣** في النماذج المعبّأة الاسترشادية.

الأربعةُ والخمسون نموذجاً المعبّأُ بلا هذا القسم كلِّه، ولا مولّدَ لها في
المستودع. فيُنسخ **جدولُ القسم من النموذج الفارغ** (بجنسه) — فهو مخرَجُ
`pbuild2` نفسُه، فيأتي بتنسيقه وكليشته وحدوده كما لو بُني معه — ويُدرج قبل
جدول «المجال ٤»، ثم تُملأ خاناتُه الثلاثُ من `m3fill.py`.

⚠️ والمقاساتُ تتبع المعبّأ لا الفارغ: التلميحُ فيه ٩ نقاط (لا ٧٫٥) والإجابةُ
   ١٠ بلون `1F4E79` — قِيس على خاناته الحرّة القائمة، وإلّا خرج القسمُ أصغرَ
   ممّا حوله فبان دخيلاً.

⚠️ ولا يُحقن مرتين: إن وُجد «المجال ٣» في المستند تُرَدّ الحقنةُ ويُقال.

الاستعمال: python3 m3inject.py [--apply] [--only "اسم الدرس"]
           بلا `--apply` يعمل على نسخةٍ في المؤقّت ولا يمسّ المسلَّم.
"""
import copy
import os
import re
import shutil
import sys

from docx import Document
from docx.shared import Pt, RGBColor

import m3fill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BASE = os.path.join(ROOT, "١ - نموذج تحضير الحصة")
MODELS = os.path.join(BASE, "نماذج معبّأة استرشادية")
BLANK = {"m": os.path.join(BASE, "بنين", "نموذج تحضير الحصة — ابن خلدون.docx"),
         "f": os.path.join(BASE, "بنات", "نموذج تحضير الحصة — ابن خلدون (بنات).docx")}
HINT_PT, ANS_PT, ANS_RGB = 9.0, 10.0, RGBColor(0x1F, 0x4E, 0x79)
ORDER = ("talk", "fair", "coop")
FLOOR = 54


def _head(tbl):
    return " ".join(tbl.rows[0].cells[0].text.split())


def _find(doc, pref):
    for t in doc.tables:
        if _head(t).startswith(pref):
            return t
    return None


def _uniq(row):
    out = []
    for c in row.cells:
        if not any(c._tc is s._tc for s in out):
            out.append(c)
    return out


def _ptext(par):
    return "".join(r.text for r in par.runs)


def _pset(par, new):
    """استبدالٌ على مستوى الفقرة: النصُّ في وورد مقطّعٌ بين runs."""
    if not par.runs:
        return
    par.runs[0].text = new
    for r in par.runs[1:]:
        r.text = ""


def fix_heads(doc, blank):
    """⛔ **العنوانُ كان يَعِد بتغطية المجال ٣** ثم صار له قسمٌ مستقلّ.
    فيُنقل عنوانُ الفارغ حرفياً — ولا يُصاغ هنا نصٌّ ثالث. وفي نسخة البنات
    كان العنوانُ القديمُ نصفَ مؤنَّث: «مشاركة المتعلمات واندماجهم»."""
    want = []
    for t in blank.tables:
        h = _head(t)
        if h.startswith("المجال ٢"):
            want.append(h)
    if len(want) != 2:
        return "لا عنوانَي «المجال ٢» في الفارغ"
    i, fixed = 0, 0
    for t in doc.tables:
        if not _head(t).startswith("المجالان ٢ و٣"):
            continue
        if i >= len(want):
            return "عناوينُ «المجالان» أكثرُ ممّا في الفارغ"
        cell = t.rows[0].cells[0]
        for par in cell.paragraphs:
            if _ptext(par).strip():
                _pset(par, want[i])
                break
        if _head(t) != want[i]:
            return "لم يُستبدل العنوان %d" % (i + 1)
        i += 1
        fixed += 1
    if fixed != 2:
        return "صُحِّح %d عنواناً لا اثنين" % fixed
    return None


def inject(src, dst, gender, key):
    txt = m3fill.TEXTS.get(key)
    if not txt:
        return "لا نصَّ لهذا الدرس في m3fill"
    doc = Document(src)
    if _find(doc, "المجال ٣") is not None:
        return "فيه «المجال ٣» أصلاً — لا يُحقن مرتين"
    after = _find(doc, "المجال ٤")
    if after is None:
        return "لا جدولَ «المجال ٤» يُدرج قبله"
    blank = Document(BLANK[gender])
    tbl = _find(blank, "المجال ٣")
    if tbl is None:
        return "لا «المجال ٣» في النموذج الفارغ — أعد بناءه"
    new = copy.deepcopy(tbl._tbl)
    after._tbl.addprevious(new)
    # الجدولُ المدرَج هو آخرُ ما أُضيف قبل «المجال ٤»
    doc_tbl = _find(doc, "المجال ٣")
    i = 0
    for ri, row in enumerate(doc_tbl.rows):
        if ri == 0:
            continue
        cells = _uniq(row)
        if len(cells) < 2:
            continue
        ans = cells[1]
        p = ans.paragraphs[0]
        for r in p.runs:
            r.font.size = Pt(HINT_PT)
        run = p.add_run("  " + txt[ORDER[i]][1 if gender == "f" else 0])
        run.font.size = Pt(ANS_PT)
        run.font.color.rgb = ANS_RGB
        base = p.runs[0]
        run.font.name = base.font.name
        run.font.bold = False
        i += 1
    if i != 3:
        return "مُلئت %d خانات لا ثلاث" % i
    err = fix_heads(doc, blank)
    if err:
        return err
    doc.save(dst)
    return None


def main():
    apply = "--apply" in sys.argv
    only = None
    if "--only" in sys.argv:
        only = sys.argv[sys.argv.index("--only") + 1]
    tmp = os.path.join("/tmp", "m3try_%d" % os.getpid())
    done, bad = 0, []
    for stage in sorted(os.listdir(MODELS)):
        sd = os.path.join(MODELS, stage)
        if not os.path.isdir(sd):
            continue
        for g, sub in (("m", "بنين"), ("f", "بنات")):
            gd = os.path.join(sd, sub)
            if not os.path.isdir(gd):
                continue
            for fn in sorted(os.listdir(gd)):
                if not fn.endswith(".docx") or fn.startswith("~$"):
                    continue
                key = re.sub(r"\s*\(بنات\)$", "", fn[:-5])
                if only and only != key:
                    continue
                src = os.path.join(gd, fn)
                if apply:
                    dst = src
                else:
                    dst = os.path.join(tmp, stage, sub, fn)
                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                err = inject(src, dst, g, key)
                if err:
                    bad.append((fn[:46], err))
                else:
                    done += 1
    for fn, e in bad:
        print("  ⛔ %-48s %s" % (fn, e))
    print("  حُقن %d نموذجاً%s" % (done, "" if apply else "  (تجربةً في %s)" % tmp))
    if not apply:
        print("  وللتطبيق على المسلَّم: python3 m3inject.py --apply")
    if done < FLOOR and not only:
        print("  ⛔ المحقونُ %d والأرضيّةُ %d" % (done, FLOOR))
        return 1
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
